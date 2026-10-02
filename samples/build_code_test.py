"""
Build samples/HBAN_FDS_Code_Test_SAMPLE.xlsx - a test workbook with ~250 FactSet codes for HBAN
and a few peers / global banks. Open it in Excel with the FactSet add-in, refresh, and run
`python tools/harvest_codes.py samples/HBAN_FDS_Code_Test_SAMPLE.xlsx` on the refreshed file
to see which codes returned data.

    python samples/build_code_test.py

Formula patterns are copied from the refreshed HBAN quarterly template (V10) and the CAMELS screen
(V5); see docs/FactSet_FDS_Coding_Guide.md.
"""
import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from tools.fds_kit import (FMT, freeze, grey_dash_rule, style_header,  # noqa: E402
                           style_period_row, style_section)

OUT = ROOT / 'samples' / 'HBAN_FDS_Code_Test_SAMPLE.xlsx'

US_PEERS = ['HBAN', 'FITB', 'RF', 'KEY', 'PNC']
GLOBAL = ['HSBA-GB', 'SAN-ES', 'NAB-AU']
LATEST_Q, PRIOR_Q = '2026/2F', '2026/1F'
HIST_Q = ['2024/3F', '2024/4F', '2025/1F', '2025/2F', '2025/3F', '2025/4F', '2026/1F', '2026/2F']
FY = 2025
EST_FY1, EST_FY2 = 2026, 2027
EST_Q = ['2026/3F', '2026/4F', '2027/1F', '2027/2F', '2027/3F', '2027/4F',
         '2028/1F', '2028/2F', '2028/3F', '2028/4F', '2029/1F', '2029/2F']

ARIAL = Font(name='Arial', size=9)
BOLD = Font(name='Arial', size=9, bold=True)
NOTE = Font(name='Arial', size=8, italic=True, color='FF595959')
INPUT_FILL = PatternFill('solid', fgColor='FFDDEBF7')
STATUS_FONT = {
    'Verified': Font(name='Arial', size=8, color='FF2E7D32'),
    'Untested': Font(name='Arial', size=8, color='FFB26A00'),
    'Not in catalog': Font(name='Arial', size=8, color='FFB00B1C'),
    'Avoid': Font(name='Arial', size=8, bold=True, color='FFB00B1C'),
}

# ------------------------------------------------------------------ catalog status lookup
CAT = {}
with open(ROOT / 'reference' / 'FactSet_Code_Catalog.csv', encoding='utf-8') as f:
    for r in csv.DictReader(f):
        v = r['Verified in a refreshed client file?']
        CAT[r['Code']] = 'Verified' if v.startswith('VERIFIED') else 'Avoid' if v.startswith('TESTED') else 'Untested'
CAT['FG_COMPANY_NAME'] = 'Verified'
CODE_RX = re.compile(r'\b((?:FFI|FF|FB|FG|FREF|P)_[A-Z0-9_]+)\s*\(')
FE_RX = re.compile(r'FE_[A-Z_]+\(\s*([A-Z0-9_]+)\s*,')


def status(fql):
    keys = CODE_RX.findall(fql) + FE_RX.findall(fql)
    found = [CAT.get(k, 'Not in catalog') for k in keys]
    for s in ('Avoid', 'Not in catalog', 'Untested'):
        if s in found:
            return s
    return 'Verified' if found else ''


def short_code(fql):
    """Readable code label for column C: drop the spliced cell references."""
    return re.sub(r'"&\$?[A-Z]+\$?\d+&"', '{p}', fql).replace('""', '"')


# ------------------------------------------------------------------ metric lists
# (category, label, fql with {p} = period cell, {q} = prior-quarter cell, units, multiplier, note)
QUARTERLY = [
    ('Valuation', 'P/E', 'FF_PE(QTR,{p})', 'PE', None, ''),
    ('Valuation', 'P/B', 'FF_PBK(QTR,{p})', 'MULT', None, ''),
    ('Valuation', 'P/TBV', 'FF_PBK_TANG(QTR,{p})', 'MULT', None, ''),
    ('Valuation', 'Market cap ($M)', 'FF_MKT_VAL(QTR,{p})', 'DOLLAR', None, ''),
    ('Valuation', 'Enterprise value ($M)', 'FF_ENTRPR_VAL(QTR,{p})', 'DOLLAR', None, ''),
    ('Valuation', 'Share price ($)', 'P_PRICE({p})', 'DOLLAR', None, ''),
    ('Shareholder returns', 'TSR - quarter (%)', 'P_PRICE_RETURNS(2,{q},{p})', 'PCT', None,
     'whole %; prior quarter -> this quarter'),
    ('Shareholder returns', 'Dividend yield (%)', 'FF_DIV_YLD(QTR,{p})', 'PCT', None, ''),
    ('Shareholder returns', 'Payout ratio (%)', 'FF_PAY_OUT_RATIO(QTR,{p})', 'PCT', None, ''),
    ('Shareholder returns', 'Dividends paid ($M)', 'FF_DIV_CF(QTR,{p})', 'DOLLAR', None, ''),
    ('Shareholder returns', 'Share repurchases ($M)', 'FF_SHS_REPURCH(TOTAL_VAL_REPURCH,QTR,{p})', 'DOLLAR', None, ''),
    ('Shareholder returns', 'DPS ($)', 'FF_DPS(QTR,{p})', 'DPS', None, ''),
    ('Shareholder returns', 'Shares outstanding (M)', 'FF_COM_SHS_OUT(QTR_R,{p})', 'NUM', None, ''),
    ('Shareholder returns', 'Stock repurchase CF ($M)', 'FF_STK_PURCH_CF(QTR_R,{p})', 'DOLLAR', None, ''),
    ('Shareholder returns', 'Company market value ($M)', 'FREF_MARKET_VALUE_COMPANY({p},,,,,0,,""LEGACY"")', 'DOLLAR', None, ''),
    ('Profitability', 'ROTCE (%)', 'FB_ROTCE(QTR,{p})', 'PCT', None, 'FF_ROTCE quarterly returned NA'),
    ('Profitability', 'ROE (%)', 'FF_ROE(QTR,{p},,,,,,""D=1"")', 'PCT', None, ''),
    ('Profitability', 'ROA (%)', 'FF_ROA(QTR_R,{p},,,,,,""D=1"")', 'PCT', None, ''),
    ('Profitability', 'EPS ($)', 'FF_EPS(QTR_R,{p})', 'DPS', None, ''),
    ('Profitability', 'Net income ($M)', 'FF_NET_INC(QTR,{p})', 'DOLLAR', None, ''),
    ('Profitability', 'PPNR ($M)', 'FF_PPNR(QTR,{p})', 'DOLLAR', None, ''),
    ('Profitability', 'Effective tax rate (%)', 'FF_TAX_RATE(QTR_R,{p})', 'PCT', None, ''),
    ('Revenue', 'Net interest income ($M)', 'FF_INT_INC_NET(QTR_R,{p})', 'DOLLAR', None, ''),
    ('Revenue', 'Non-interest income ($M)', 'FF_NON_INT_INC(QTR,{p})', 'DOLLAR', None, ''),
    ('Revenue', 'Total revenue ($M)', '=SUM:Net interest income ($M)|Non-interest income ($M)', 'DOLLAR', None,
     'cell math: NII + non-interest income'),
    ('Efficiency', 'Efficiency ratio (%) - FF_EFF_RATIO x100', 'FF_EFF_RATIO(QTR_R,{p})', 'PCT', 100,
     'code returns a fraction, x100'),
    ('Efficiency', 'Efficiency ratio (%) - FF_BK_EFF_RATIO', 'FF_BK_EFF_RATIO(QTR,{p})', 'PCT', None,
     'whole %; compare with the row above'),
    ('Efficiency', 'Non-interest expense ($M)', 'FF_NON_INT_EXP(QTR,{p})', 'DOLLAR', None, ''),
    ('Efficiency', 'Personnel expense ($M)', 'FF_LABOR_EXP(QTR,{p})', 'DOLLAR', None, ''),
    ('NIM & funding', 'NIM (%)', 'FB_INT_MGN(QTR,{p})', 'PCT', None, 'FF_INT_MGN quarterly returned NA for HBAN'),
    ('NIM & funding', 'Earning-asset yield (%)', 'FF_YLD_INT_EARN_ASSETS(QTR_R,{p})', 'PCT', None, ''),
    ('NIM & funding', 'Total interest expense ($M)', 'FF_INT_EXP_TOT(QTR,{p})', 'DOLLAR', None, ''),
    ('NIM & funding', 'Avg interest-bearing liabilities ($M)', 'FF_BK_AVG_LIABS_INTB(QTR_TOT,{p})', 'DOLLAR', None, ''),
    ('NIM & funding', 'Cost of funds (%, annualized)',
     '=COF:Total interest expense ($M)|Avg interest-bearing liabilities ($M)', 'PCT', None,
     'cell math: int. expense / avg IBL x100 x4'),
    ('NIM & funding', 'Deposit cost (%, annualized)', 'FF_COST_DEPS(QTR,{p})', 'PCT', 4, 'quarterly rate, x4'),
    ('NIM & funding', 'IB liability cost (%)', 'FF_BK_INT_COST_INTB_AVG(QTR,{p})', 'PCT', None,
     'matched cost of funds in an earlier test'),
    ('NIM & funding', 'Funding cost (%)', 'FF_BK_COST(QTR_NINTB_LIABS,{p})', 'PCT', None,
     'matched cost of funds in an earlier test'),
    ('NIM & funding', 'Net interest spread (%)', 'FF_BK_INT_SPREAD(QTR,{p})', 'PCT', None, ''),
    ('NIM & funding', 'Tax-equivalent NIM (%)', 'FF_BK_INT_MGN_TAX_EQV(QTR,{p})', 'PCT', None, ''),
    ('Balance sheet', 'Total assets ($M)', 'FF_ASSETS(QTR,{p})', 'DOLLAR', None, ''),
    ('Balance sheet', 'Total liabilities ($M)', 'FF_LIABS(QTR_R,{p})', 'DOLLAR', None, ''),
    ('Balance sheet', 'Intangible assets ($M)', 'FF_INTANG(QTR_R,{p})', 'DOLLAR', None, ''),
    ('Balance sheet', 'Average assets ($M)', 'FF_BK_AVG_ASSETS(QTR,{p})', 'DOLLAR', None, ''),
    ('Balance sheet', 'Avg earning assets ($M)', 'FF_BK_AVG_ASSETS(QTR_INTE,{p})', 'DOLLAR', None, 'an average, not EOP'),
    ('Balance sheet', 'Total loans ($M)', 'FF_BK_LOAN_TOT(QTR,{p})', 'DOLLAR', None, ''),
    ('Balance sheet', 'Total deposits ($M)', 'FF_DEPS(QTR,{p})', 'DOLLAR', None, ''),
    ('Balance sheet', 'Interest-bearing deposits ($M)', 'FF_DEPS_INTB(QTR,{p})', 'DOLLAR', None, ''),
    ('Balance sheet', 'Non-interest-bearing deposits ($M)', 'FF_DEPS_NINTB(QTR,{p})', 'DOLLAR', None, ''),
    ('Balance sheet', 'Loans / deposits (%)', '=LTD:Total loans ($M)|Total deposits ($M)', 'PCT', None,
     'cell math: loans / deposits x100'),
    ('Balance sheet', 'Tangible common equity ($M)', 'FF_COM_EQ_TANG(QTR,{p})', 'DOLLAR', None, ''),
    ('Balance sheet', 'Avg tangible common equity ($M)', 'FF_BK_SUPPL_AVG(QTR_TCE,{p})', 'DOLLAR', None, ''),
    ('Balance sheet', 'TBVPS ($)', 'FF_BPS_TANG(QTR,{p})', 'DPS', None, ''),
    ('Credit', 'NCO ratio (%, annualized)', 'FF_CHARGE_OFFS_LOANS_PCT(QTR_R,{p})', 'PCT', 4, 'quarterly rate, x4'),
    ('Credit', 'NPL ratio (%)', 'FF_NONPERF_LOAN_PCT(QTR_R,{p})', 'PCT', None, ''),
    ('Credit', 'NPA / assets (%)', 'FF_NPA_ASSETS_PCT(QTR,{p})', 'PCT', None, ''),
    ('Capital', 'CET1 ratio (%)', 'FF_BK_COM_EQ_TIER1_RATIO(QTR,{p})', 'PCT', None, ''),
    ('Capital', 'Risk-weighted assets ($M)', 'FF_ASSETS_RISK_WGHT(QTR,{p})', 'DOLLAR', None, ''),
    ('Franchise', 'Branch count', 'FB_DOM_OFFCE_NUM(QTR,{p})', 'COUNT', None, ''),
    ('Consensus (history)', 'Consensus EPS - quarter ($)', "FE_ESTIMATE(EPS_NONGAAP,MEAN,QTR,{p},0,,,'')", 'DPS', None, ''),
    ('Consensus (history)', 'Consensus revenue - quarter ($M)', "FE_ESTIMATE(SALES,MEAN,QTR_ROLL,{p},0,,,'')", 'DOLLAR', None, ''),
    ('Consensus (history)', 'Consensus NII - quarter ($M)', 'FE_ESTIMATE(INT_INC_NET,MEAN,QTR,{p})', 'DOLLAR', None, ''),
    ('Consensus (history)', 'Consensus fee income - quarter ($M)', "FE_ESTIMATE(INCFEESCOM,MEAN,QTR_ROLL,{p},0,,,'')", 'DOLLAR', None, ''),
    ('Consensus (history)', 'Consensus ROTCE - quarter (%)', "FE_ESTIMATE(ROTE,MEAN,QTR_ROLL,{p},0,,,'')", 'PCT', None, ''),
    ('Consensus (history)', 'Consensus CET1 capital - quarter ($M)', "FE_ESTIMATE(COM_EQUITY_TIER1,MEAN,QTR,{p},,,,'')", 'DOLLAR', None, ''),
]

# CAMELS annual: (category, metric, FFI, FF, FB, units).  {y} = fiscal-year cell.
CAMELS = [
    ('Capital', 'CET1 capital ($M)', 'FFI_COM_EQ_TIER1_TOT(ANN_L,{y},,,,USD)', 'FF_BK_COM_EQ_TIER1_TOT(ANN_R,{y},,,,USD)',
     'FB_COM_EQ_TIER1(ANN,{y},,,RF,USD)', 'DOLLAR'),
    ('Capital', 'CET1 ratio (%)', 'FFI_COM_EQ_TIER1_RATIO(ANN_L,{y})', 'FF_BK_COM_EQ_TIER1_RATIO(ANN_R,{y})',
     'FB_COM_EQ_TIER1_RATIO(ANN,{y},,,RF)', 'PCT'),
    ('Capital', 'Tier 1 capital ($M)', 'FFI_TIER1_CAP(ANN_L,{y},,,,USD)', 'FF_TIER1_CAP(ANN_R,{y},,,,USD)',
     'FB_TIER1_CAP(ANN,{y},,,RF,USD)', 'DOLLAR'),
    ('Capital', 'Tier 1 ratio (%)', 'FFI_CAP_RATIO_TIER1(ANN_L,{y})', 'FF_CAP_RATIO_TIER1(ANN_R,{y})',
     'FB_CAP_RATIO_TIER1(ANN,{y},,,RF)', 'PCT'),
    ('Capital', 'Total capital ratio (%)', 'FFI_CAP_RATIO_TOT(ANN_L,{y})', 'FF_CAP_RATIO_TOT(ANN_R,{y})',
     'FB_CAP_RATIO_TOT(ANN,{y},,,RF)', 'PCT'),
    ('Capital', 'Leverage ratio (%)', 'FFI_LEV_RATIO_RPT(ANN_L,{y})', 'FF_BK_LEV_RATIO(ANN_R,{y})',
     'FB_LEV_RATIO(ANN,{y},,,RF)', 'PCT'),
    ('Capital', 'RWA ($M)', 'FFI_RWA(ANN_L,{y},,,,USD)', 'FF_ASSETS_RISK_WGHT(ANN_R,{y},,,,USD)',
     'FB_ASSETS_RISK_WGHT(ANN,{y},,,RF,USD)', 'DOLLAR'),
    ('Asset quality', 'Non-performing loans ($M)', 'FFI_NPL_LOAN_ADV(ANN_L,{y},,,,USD)', 'FF_LOAN_NONPERF(ANN_R,{y},,,,USD)',
     'FB_NONPERF_LOAN(ANN,{y},,,RF,USD)', 'DOLLAR'),
    ('Asset quality', 'Provision for credit losses ($M)', 'FFI_LOAN_LOSS_PROV(ANN_L,{y},,,,USD)',
     'FF_LOAN_LOSS_PROV(ANN_R,{y},,,,USD)', 'FB_LOAN_LOSS_PROV(ANN,{y},,,RF,USD)', 'DOLLAR'),
    ('Asset quality', 'Net charge-offs ($M)', 'FFI_CHRG_OFF_NET(ANN_L,{y},,,,USD)', None, None, 'DOLLAR'),
    ('Asset quality', 'NPLs / allowance (%)', 'FFI_NPL_LOAN_LOSS_RSRV_RATIO(ANN_L,{y})', 'FF_NONPERF_LOAN_LOSS_RSRV(ANN_R,{y})',
     None, 'PCT'),
    ('Asset quality', 'NPLs / loans (%)', 'FFI_NPL_LOAN_RATIO(ANN_L,{y})', 'FF_NONPERF_LOAN_PCT(ANN_R,{y})',
     'FB_NONPERF_LOAN_PCT(ANN,{y},,,RF)', 'PCT'),
    ('Management', 'Efficiency ratio (%)', 'FFI_EFF_RATIO(ANN_L,{y})', 'FF_BK_EFF_RATIO(ANN_R,{y})',
     'FB_EFF_RATIO(ANN,{y},,,RF)', 'PCT'),
    ('Earnings', 'ROTE (%)', 'FFI_ROTE(ANN_L,{y})', None, 'FB_ROTE(ANN,{y},,,RF)', 'PCT'),
    ('Earnings', 'ROTCE (%)', 'FFI_ROTCE(ANN_L,{y})', 'FF_ROTCE(ANN_R,{y})', 'FB_ROTCE(ANN,{y},,,RF)', 'PCT'),
    ('Earnings', 'ROA (%)', 'FFI_ROA(ANN_L,{y})', 'FF_ROA(ANN_R,{y})', 'FB_ROA(ANN,{y},,,RF)', 'PCT'),
    ('Earnings', 'NIM (%)', 'FFI_AVG_BAL_INT_RATE_NET_MGN(ANN_L,{y})', 'FF_INT_MGN(ANN_R,{y})',
     'FB_INT_MGN(ANN,{y},,,RF)', 'PCT'),
    ('Earnings', 'Net interest income ($M)', 'FFI_INT_INC_NET(ANN_L,{y},,,,USD)', 'FF_INT_INC_NET(ANN_R,{y},,,,USD)',
     'FB_INT_INC_NET(ANN,{y},,,RF,USD)', 'DOLLAR'),
    ('Earnings', 'Non-interest income ($M)', 'FFI_NON_INT_INC(ANN_L,{y},,,,USD)', 'FF_NON_INT_INC(ANN_R,{y},,,,USD)',
     'FB_NON_INT_INC(ANN,{y},,,RF,USD)', 'DOLLAR'),
    ('Liquidity', 'Loans ($M)', 'FFI_LOAN_ADV_TOT(ANN_L,{y},,,,USD)', 'FF_BK_LOAN_TOT(ANN_R,{y},,,,USD)',
     'FB_TOT_HFI_HFS_UNEARN_INC(ANN,{y},,,RF,USD)', 'DOLLAR'),
    ('Liquidity', 'Deposits ($M)', 'FFI_DEPS_TOT(ANN_L,{y},,,,USD)', 'FF_DEPS(ANN_R,{y},,,,USD)',
     'FB_DEPS(ANN,{y},,,RF,USD)', 'DOLLAR'),
    ('Liquidity', 'Liquidity coverage ratio (%)', 'FFI_BK_LIQ_COVG_RATIO(ANN_L,{y})', 'FF_BK_LIQ_COVG_RATIO(ANN_R,{y})',
     None, 'PCT'),
    ('Liquidity', 'Net stable funding ratio (%)', 'FFI_BK_NSFR(ANN_L,{y})', None, None, 'PCT'),
    ('Sensitivity', 'Loans to customers @ amortized cost ($M)', 'FFI_LOAN_AMORT_CUST(ANN_L,{y},,,,USD)',
     'FF_LOAN_NET(ANN_R,{y},,,,USD)', None, 'DOLLAR'),
    ('Sensitivity', 'Loans to FIs @ amortized cost ($M)', 'FFI_LOAN_AMORT_BK(ANN_L,{y},,,,USD)', None, None, 'DOLLAR'),
]

# Estimates: (category, label, fql, units). {f1}/{f2} = fiscal-year cells.
NTM = "FE_ESTIMATE({i},MEAN,NTMA,,NOW,,,'CURRENCY=USD')"
FYE = "FE_ESTIMATE({i},MEAN,ANN_ROLL,{f},NOW,,,'CURRENCY=USD')"
EST_ITEMS = [
    ('EPS', 'EPS ($)', 'DPS'), ('SALES', 'Revenue ($M)', 'DOLLAR'),
    ('NETINTERESTINC', 'Net interest income ($M)', 'DOLLAR'), ('NON_INT_INC', 'Non-interest income ($M)', 'DOLLAR'),
    ('INT_INC_MARGIN', 'NIM (%)', 'PCT'), ('COST_INCOME', 'Efficiency ratio (%)', 'PCT'),
    ('OPEREXPEN', 'Non-interest expense ($M)', 'DOLLAR'), ('EBIT_ADJ', 'PPNR ($M)', 'DOLLAR'),
    ('NETPROFIT', 'Net income ($M)', 'DOLLAR'), ('ROTE', 'ROTCE (%)', 'PCT'),
    ('COMCAP_RATIO_TIER1', 'CET1 ratio (%)', 'PCT'), ('LOAN_GROSS', 'Gross loans ($M)', 'DOLLAR'),
    ('DEPS', 'Deposits ($M)', 'DOLLAR'), ('NET_CHARGE_OFFS', 'Net charge-offs ($M)', 'DOLLAR'),
    ('NETCHARGE_LOANNET', 'NCO / loans (%, annualized)', 'PCT'), ('LOAN_PROV', 'Loan-loss provision ($M)', 'DOLLAR'),
    ('NETDIV', 'DPS ($)', 'DPS'), ('BPS_TANG', 'TBVPS ($)', 'DPS'),
]
VALUATION = [('PE', 'P/E', 'PE'), ('PBPS', 'P/B', 'MULT'), ('P_TBPS', 'P/TBV', 'MULT'), ('DIV_YLD', 'Dividend yield (%)', 'PCT')]

# HBAN FE_TIMESERIES spills: (category, label, item, valuation?, units)
TS = [
    ('Valuation', 'P/E', 'PE', True, 'PE'), ('Valuation', 'P/B', 'PBPS', True, 'MULT'),
    ('Valuation', 'P/TBV', 'P_TBPS', True, 'MULT'), ('Valuation', 'Dividend yield (%)', 'DIV_YLD', True, 'PCT'),
    ('Profitability', 'ROTCE (%)', 'ROTE', False, 'PCT'), ('Profitability', 'ROE (%)', 'ROE', True, 'PCT'),
    ('Profitability', 'ROA (%)', 'ROA', True, 'PCT'), ('Profitability', 'EPS ($)', 'EPS', False, 'DPS'),
    ('Profitability', 'EPS - custom ($)', 'CUSTOM_EPS', False, 'DPS'),
    ('Profitability', 'Net income ($M)', 'NETPROFIT', False, 'DOLLAR'), ('Profitability', 'PPNR ($M)', 'EBIT_ADJ', False, 'DOLLAR'),
    ('Revenue', 'Net interest income ($M)', 'NETINTERESTINC', False, 'DOLLAR'),
    ('Revenue', 'Non-interest income ($M)', 'NON_INT_INC', False, 'DOLLAR'),
    ('Revenue', 'Revenue ($M)', 'SALES', False, 'DOLLAR'), ('Revenue', 'Fee income ($M)', 'INCFEESCOM', False, 'DOLLAR'),
    ('Efficiency', 'Efficiency ratio (%)', 'COST_INCOME', False, 'PCT'),
    ('Efficiency', 'Non-interest expense ($M)', 'OPEREXPEN', False, 'DOLLAR'),
    ('Efficiency', 'Personnel expense ($M)', 'SAL_BENEFITS', False, 'DOLLAR'),
    ('NIM', 'NIM (%)', 'INT_INC_MARGIN', False, 'PCT'),
    ('Balance sheet', 'Total assets ($M)', 'TOTASSET', False, 'DOLLAR'),
    ('Balance sheet', 'Avg earning assets ($M)', 'AVG_EARN_ASSETS', False, 'DOLLAR'),
    ('Balance sheet', 'Gross loans ($M)', 'LOAN_GROSS', False, 'DOLLAR'),
    ('Balance sheet', 'Deposits ($M)', 'DEPS', False, 'DOLLAR'), ('Balance sheet', 'TBVPS ($)', 'BPS_TANG', False, 'DPS'),
    ('Balance sheet', 'Intangibles ($M)', 'INTANG', False, 'DOLLAR'),
    ('Credit', 'NCO / loans (%)', 'NETCHARGE_LOANNET', True, 'PCT'), ('Credit', 'NPL ratio (%)', 'NPL', True, 'PCT'),
    ('Credit', 'NPAs / loans + OREO (%)', 'LOANOREO_NONPERF', False, 'PCT'),
    ('Capital', 'CET1 ratio (%)', 'COMCAP_RATIO_TIER1', False, 'PCT'),
    ('Capital', 'CET1 capital ($M)', 'COM_EQUITY_TIER1', False, 'DOLLAR'),
    ('Capital', 'RWA ($M)', 'ASSETS_RISK_WGHT', False, 'DOLLAR'),
    ('Shareholder returns', 'Share repurchases ($M)', 'SHS_REPURCH', False, 'DOLLAR'),
    ('Shareholder returns', 'DPS ($)', 'NETDIV', False, 'DPS'),
]


# ------------------------------------------------------------------ helpers
def splice(fql, **cells):
    for k, cell in cells.items():
        fql = fql.replace('{' + k + '}', f'"&{cell}&"')
    return fql


def pull(id_cell, fql, mult=None):
    inner = f'FDS({id_cell},"{fql}")' + (f'*{mult}' if mult else '')
    return f'=_xlfn.IFNA({inner},"-")'


def label_cols(ws, row, cat, label, code, st):
    ws.cell(row, 1, cat).font = NOTE
    ws.cell(row, 2, label).font = ARIAL
    ws.cell(row, 3, code).font = Font(name='Consolas', size=8)
    c = ws.cell(row, 4, st)
    c.font = STATUS_FONT.get(st, NOTE)


def header_block(ws, title, note, inputs, id_row, ids, first_col=5):
    ws['A1'] = title
    ws['A1'].font = Font(name='Arial', size=12, bold=True, color='FF1F4E78')
    ws['A2'] = note
    ws['A2'].font = NOTE
    for i, (lab, val) in enumerate(inputs):
        ws.cell(3 + i, 2, lab).font = BOLD
        c = ws.cell(3 + i, 3, val)
        c.fill, c.font = INPUT_FILL, BOLD
        c.alignment = Alignment(horizontal='center')
    for col, txt in enumerate(['Category', 'Metric', 'FactSet code', 'Catalog status'], 1):
        ws.cell(id_row, col, txt)
    for j, t in enumerate(ids):
        ws.cell(id_row, first_col + j, t)
    style_header([ws.cell(id_row, c) for c in range(1, first_col + len(ids))])
    for c in range(first_col, first_col + len(ids)):
        ws.cell(id_row, c).fill = INPUT_FILL


def widths(ws, ncols, first=5, w=14):
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 36
    ws.column_dimensions['C'].width = 46
    ws.column_dimensions['D'].width = 12
    for c in range(first, first + ncols):
        ws.column_dimensions[get_column_letter(c)].width = w


# ------------------------------------------------------------------ sheets
wb = Workbook()
readme = wb.active
readme.title = 'Read Me'


def peers_quarter_sheet():
    ws = wb.create_sheet('Quarterly - Peers')
    ID_ROW, NAME_ROW, FIRST = 6, 7, 8
    header_block(ws, 'Quarterly code test - US peers, one quarter',
                 '$ in millions (USD), per-share in $, % as whole numbers. Blue cells are inputs - change the '
                 'quarter or tickers and refresh.',
                 [('Quarter', LATEST_Q), ('Prior quarter (TSR start)', PRIOR_Q)], ID_ROW, US_PEERS)
    label_cols(ws, NAME_ROW, '', 'Company name', 'FG_COMPANY_NAME', 'Verified')
    for j in range(len(US_PEERS)):
        col = get_column_letter(5 + j)
        c = ws.cell(NAME_ROW, 5 + j, f'=_xlfn.IFNA(FDS({col}${ID_ROW},"FG_COMPANY_NAME"),"-")')
        c.font, c.alignment = NOTE, Alignment(horizontal='center', wrap_text=True)
    write_metric_rows(ws, FIRST, ID_ROW, lambda col: ('$C$3', '$C$4'), len(US_PEERS))
    freeze(ws, f'E{FIRST}')
    widths(ws, len(US_PEERS))


def hban_history_sheet():
    ws = wb.create_sheet('HBAN Quarterly')
    ws['A1'] = 'Quarterly code test - HBAN, 8 quarters'
    ws['A1'].font = Font(name='Arial', size=12, bold=True, color='FF1F4E78')
    ws['A2'] = ('$ in millions (USD), per-share in $, % as whole numbers. Blue cells are inputs: ticker in C3, '
                'quarters in row 5 (row 6 = prior quarter, used for TSR).')
    ws['A2'].font = NOTE
    ws['B3'] = 'Ticker'
    ws['B3'].font = BOLD
    ws['C3'] = 'HBAN'
    ws['C3'].fill, ws['C3'].font = INPUT_FILL, BOLD
    PROW, QROW, FIRST = 5, 6, 8
    for col, txt in enumerate(['Category', 'Metric', 'FactSet code', 'Catalog status'], 1):
        ws.cell(PROW, col, txt)
    prior = ['2024/2F'] + HIST_Q[:-1]
    for j, (p, q) in enumerate(zip(HIST_Q, prior)):
        ws.cell(PROW, 5 + j, p).fill = INPUT_FILL
        ws.cell(QROW, 5 + j, q)
    style_header([ws.cell(PROW, c) for c in range(1, 5 + len(HIST_Q))])
    for j in range(len(HIST_Q)):
        ws.cell(PROW, 5 + j).fill = INPUT_FILL
    ws.cell(QROW, 2, 'Prior quarter (TSR start)')
    style_period_row([ws.cell(QROW, c) for c in range(2, 5 + len(HIST_Q))])
    write_metric_rows(ws, FIRST, None, lambda col: (f'{col}${PROW}', f'{col}${QROW}'), len(HIST_Q), id_cell='$C$3')
    freeze(ws, f'E{FIRST}')
    widths(ws, len(HIST_Q), w=12)


def write_metric_rows(ws, first_row, id_row, periods, ncols, id_cell=None):
    row, rows_by_label, last_cat = first_row, {}, None
    for cat, label, fql, units, mult, note in QUARTERLY:
        if cat != last_cat:
            ws.cell(row, 1, cat)
            style_section([ws.cell(row, c) for c in range(1, 5 + ncols)])
            row += 1
            last_cat = cat
        calc = fql.startswith('=')
        code = note if calc else short_code(fql) + (f'  x{mult}' if mult else '')
        label_cols(ws, row, '', label, code, 'Cell math' if calc else status(fql))
        if note and not calc:
            ws.cell(row, 3).value = f'{code}   ({note})'
        for j in range(ncols):
            col = get_column_letter(5 + j)
            idc = id_cell or f'{col}${id_row}'
            p, q = periods(col)
            if calc:
                kind, refs = fql[1:].split(':')
                a, b = (f'{col}{rows_by_label[x]}' for x in refs.split('|'))
                f = {'SUM': f'=IFERROR({a}+{b},"-")',
                     'COF': f'=IFERROR(({a}/{b})*100*4,"-")',
                     'LTD': f'=IFERROR(({a}/{b})*100,"-")'}[kind]
            else:
                f = pull(idc, splice(fql, p=p, q=q), mult)
            c = ws.cell(row, 5 + j, f)
            c.number_format, c.font = FMT[units], ARIAL
        rows_by_label[label] = row
        row += 1
    grey_dash_rule(ws, f'E{first_row}:{get_column_letter(4 + ncols)}{row}')


def camels_sheet():
    ids = US_PEERS[:3] + GLOBAL
    ws = wb.create_sheet('Annual - CAMELS Test')
    ID_ROW, NAME_ROW, FIRST = 6, 7, 8
    header_block(ws, 'Annual code test - FFI vs FF vs FB, and the @ fallback chain',
                 'Each metric is pulled four ways: FFI_, FF_, FB_ (US banks only) and the FFI@FF@FB chain. '
                 'This shows which family works per bank and whether the cross-family chain fills in. '
                 '$ in millions USD; % as whole numbers.',
                 [('Fiscal year', FY)], ID_ROW, ids)
    label_cols(ws, NAME_ROW, '', 'Company name', 'FG_COMPANY_NAME', 'Verified')
    for j in range(len(ids)):
        col = get_column_letter(5 + j)
        c = ws.cell(NAME_ROW, 5 + j, f'=FDSC("-",{col}${ID_ROW},"FG_COMPANY_NAME")')
        c.font, c.alignment = NOTE, Alignment(horizontal='center', wrap_text=True)
    row, last_cat = FIRST, None
    for cat, metric, ffi_, ff_, fb_, units in CAMELS:
        if cat != last_cat:
            ws.cell(row, 1, cat)
            style_section([ws.cell(row, c) for c in range(1, 5 + len(ids))])
            row += 1
            last_cat = cat
        variants = [(f'{metric} - FFI', ffi_), (f'{metric} - FF', ff_), (f'{metric} - FB (US only)', fb_)]
        variants = [(lab, x) for lab, x in variants if x]
        if len(variants) > 1:
            variants.append((f'{metric} - CHAIN (FFI@FF@FB)', '@'.join(x for _, x in variants)))
        for k, (lab, fql) in enumerate(variants):
            is_chain = 'CHAIN' in lab
            label_cols(ws, row, '', lab, short_code(fql), 'Unproven chain' if is_chain else status(fql))
            if is_chain:
                ws.cell(row, 4).font = STATUS_FONT['Untested']
                ws.cell(row, 2).font = BOLD
            for j in range(len(ids)):
                col = get_column_letter(5 + j)
                c = ws.cell(row, 5 + j, f'=FDSC("-",{col}${ID_ROW},"{splice(fql, y="$C$3")}")')
                c.number_format, c.font = FMT[units], BOLD if is_chain else ARIAL
            row += 1
    grey_dash_rule(ws, f'E{FIRST}:{get_column_letter(4 + len(ids))}{row}')
    freeze(ws, f'E{FIRST}')
    widths(ws, len(ids))
    ws.column_dimensions['B'].width = 48
    ws.column_dimensions['C'].width = 60


def estimates_sheet():
    ws = wb.create_sheet('Estimates - Peers')
    ID_ROW, NAME_ROW, FIRST = 7, 8, 9
    header_block(ws, 'Consensus estimate code test - US peers',
                 'Mean consensus (FE_). NTM = next twelve months; FY = fiscal year in C3 / C4. '
                 '$ in millions USD; % as whole numbers.',
                 [('Fiscal year 1', EST_FY1), ('Fiscal year 2', EST_FY2)], ID_ROW, US_PEERS)
    label_cols(ws, NAME_ROW, '', 'Company name', 'FG_COMPANY_NAME', 'Verified')
    for j in range(len(US_PEERS)):
        col = get_column_letter(5 + j)
        c = ws.cell(NAME_ROW, 5 + j, f'=_xlfn.IFNA(FDS({col}${ID_ROW},"FG_COMPANY_NAME"),"-")')
        c.font, c.alignment = NOTE, Alignment(horizontal='center', wrap_text=True)
    blocks = [('NTM (next twelve months)', lambda i: NTM.format(i=i)),
              ('Fiscal year 1 (C3)', lambda i: FYE.format(i=i, f='{f1}')),
              ('Fiscal year 2 (C4)', lambda i: FYE.format(i=i, f='{f2}'))]
    rows = []
    for bname, fn in blocks:
        rows.append(('section', bname))
        for item, lab, units in EST_ITEMS:
            rows.append((bname.split(' (')[0], lab, fn(item), units))
    rows.append(('section', 'Next fiscal year valuation (FE_VALUATION)'))
    for item, lab, units in VALUATION:
        rows.append(('Valuation', lab, f"FE_VALUATION({item},MEAN,ANN_ROLL,+1,0,,,'')", units))
    rows.append(('section', 'Next quarter'))
    rows.append(('Next quarter', 'EPS - next quarter ($)', "FE_ESTIMATE(EPS,MEAN,QTR_ROLL,+1,0,,,'')", 'DPS'))
    rows.append(('Next quarter', 'Revenue - next quarter ($M)', "FE_ESTIMATE(SALES,MEAN,QTR_ROLL,+1,0,,,'')", 'DOLLAR'))
    rows.append(('Next quarter', 'NIM - next quarter (%)', "FE_ESTIMATE(INT_INC_MARGIN,MEAN,QTR_ROLL,+1,0,,,'')", 'PCT'))
    rows.append(('section', 'LTM actuals from the estimates database'))
    for item, lab, units in [('EPS', 'EPS - LTM actual ($)', 'DPS'), ('SALES', 'Revenue - LTM actual ($M)', 'DOLLAR'),
                             ('LOAN_PROV', 'Loan-loss provision - LTM actual ($M)', 'DOLLAR')]:
        rows.append(('LTM actual', lab, f"FE_ESTIMATE({item},MEAN,LTMA,,NOW,,,'CURRENCY=USD')", units))
    row = FIRST
    for r in rows:
        if r[0] == 'section':
            ws.cell(row, 1, r[1])
            style_section([ws.cell(row, c) for c in range(1, 5 + len(US_PEERS))])
            row += 1
            continue
        _, lab, fql, units = r
        label_cols(ws, row, '', lab, short_code(fql), status(fql))
        for j in range(len(US_PEERS)):
            col = get_column_letter(5 + j)
            c = ws.cell(row, 5 + j, pull(f'{col}${ID_ROW}', splice(fql, f1='$C$3', f2='$C$4')))
            c.number_format, c.font = FMT[units], ARIAL
        row += 1
    grey_dash_rule(ws, f'E{FIRST}:{get_column_letter(4 + len(US_PEERS))}{row}')
    freeze(ws, f'E{FIRST}')
    widths(ws, len(US_PEERS))
    ws.column_dimensions['C'].width = 58


def timeseries_sheet():
    ws = wb.create_sheet('HBAN Estimates TS')
    ws['A1'] = 'Consensus time-series test - HBAN, 12 future quarters (FE_TIMESERIES spills)'
    ws['A1'].font = Font(name='Arial', size=12, bold=True, color='FF1F4E78')
    ws['A2'] = ('Each row is ONE formula in column E that spills right across the 12 quarters - leave F:P empty. '
                'Start quarter = E5, end quarter = P5. Currency = as reported (USD for HBAN).')
    ws['A2'].font = NOTE
    ws['B3'], ws['C3'] = 'Ticker', 'HBAN'
    ws['B3'].font = BOLD
    ws['C3'].fill, ws['C3'].font = INPUT_FILL, BOLD
    PROW, FIRST = 5, 7
    for col, txt in enumerate(['Category', 'Metric', 'FE item', 'Catalog status'], 1):
        ws.cell(PROW, col, txt)
    for j, p in enumerate(EST_Q):
        ws.cell(PROW, 5 + j, p)
    style_header([ws.cell(PROW, c) for c in range(1, 5 + len(EST_Q))])
    for j in range(len(EST_Q)):
        ws.cell(PROW, 5 + j).fill = INPUT_FILL
    last = get_column_letter(4 + len(EST_Q))
    row, last_cat = FIRST, None
    for cat, lab, item, val, units in TS:
        if cat != last_cat:
            ws.cell(row, 1, cat)
            style_section([ws.cell(row, c) for c in range(1, 5 + len(EST_Q))])
            row += 1
            last_cat = cat
        fn = 'FE_TIMESERIES_VALUATION' if val else 'FE_TIMESERIES'
        opts = ("'BKRACTMED=1,WIN=0,UNITS=AUTO,DATE=NOW,CALC=LTMA'" if val
                else "'BKRACTMED=1,WIN=0,CURRENCY=RPT,UNITS=AUTO,DATE=NOW'")
        fql = f'{fn}({item},MEAN,"&$E${PROW}&","&${last}${PROW}&",FQ,{opts})'
        label_cols(ws, row, '', lab, f'{fn}({item})', status(f'{fn}({item},'))
        ws.cell(row, 5, f'=FDSRC("-",$C$3,"{fql}")').font = ARIAL
        for j in range(len(EST_Q)):
            ws.cell(row, 5 + j).number_format = FMT[units]
            ws.cell(row, 5 + j).font = ARIAL
        row += 1
    # guidance example
    ws.cell(row, 1, 'Guidance')
    style_section([ws.cell(row, c) for c in range(1, 5 + len(EST_Q))])
    row += 1
    fql = (f'FE_TIMESERIES_GUIDANCE(COST_INCOME,LOW,"&$E${PROW}&","&${last}${PROW}&",FQ,'
           f"'BKRACTMED=1,WIN=0,CURRENCY=RPT,UNITS=AUTO,DATE=NOW')")
    label_cols(ws, row, '', 'Efficiency ratio guidance - low (%)', 'FE_TIMESERIES_GUIDANCE(COST_INCOME,LOW)',
               'Untested')
    ws.cell(row, 5, f'=FDSRC("-",$C$3,"{fql}")').font = ARIAL
    for j in range(len(EST_Q)):
        ws.cell(row, 5 + j).number_format = FMT['PCT']
    grey_dash_rule(ws, f'E{FIRST}:{last}{row}')
    freeze(ws, f'E{FIRST}')
    widths(ws, len(EST_Q), w=11)
    ws.column_dimensions['C'].width = 40


def write_readme():
    ws = readme
    lines = [
        ('FactSet =FDS code test workbook', 'title'),
        ('Purpose: check which FactSet codes return data for HBAN, US peers and a few global banks.', ''),
        ('', ''),
        ('How to use', 'h'),
        ('1. Open in Excel with the FactSet add-in loaded, then refresh (FactSet ribbon > Refresh / Recalculate).', ''),
        ('2. Blue cells are inputs - tickers, quarters and fiscal years can be changed.', ''),
        ('3. "-" means no data for that code / company / period (shown in grey).', ''),
        ('4. Save the refreshed file and send it back - Claude runs tools/harvest_codes.py on it to mark which', ''),
        ('   codes returned numbers and updates the catalog.', ''),
        ('', ''),
        ('Sheets', 'h'),
        ('Quarterly - Peers      ~65 quarterly codes for HBAN, FITB, RF, KEY, PNC in one quarter (C3).', ''),
        ('HBAN Quarterly         the same codes for HBAN across 8 quarters (2024/3F - 2026/2F).', ''),
        ('Annual - CAMELS Test   25 CAMELS metrics pulled via FFI_, FF_, FB_ and the FFI@FF@FB chain, for 3 US and', ''),
        ('                       3 global banks (HSBA-GB, SAN-ES, NAB-AU). Tests whether cross-family chains work.', ''),
        ('Estimates - Peers      consensus NTM / FY1 / FY2, valuation, next quarter and LTM actuals (FE_).', ''),
        ('HBAN Estimates TS      FE_TIMESERIES spills for 12 future quarters (one formula per row in column E).', ''),
        ('', ''),
        ('Catalog status column', 'h'),
        ('Verified        the code returned numbers in an earlier refreshed file.', ''),
        ('Untested        the code is in FactSet\'s code lists but has not been tested yet.', ''),
        ('Not in catalog  the code is not in FactSet\'s code lists - most likely to fail.', ''),
        ('Unproven chain  cross-family @ fallback (FFI@FF@FB) - check it fills in where a single family works.', ''),
        ('Cell math       calculated in Excel from the rows above.', ''),
        ('', ''),
        ('Units', 'h'),
        ('$ in millions; per-share in $; % as whole numbers (10.5 = 10.5%). NCO ratio and deposit cost are', ''),
        ('quarterly rates x4; FF_EFF_RATIO is a fraction x100. FB_ codes cover US banks only.', ''),
        ('Outside Excel+FactSet every formula shows #NAME? - that is expected.', ''),
    ]
    for i, (t, kind) in enumerate(lines, 1):
        c = ws.cell(i, 1, t)
        c.font = (Font(name='Arial', size=14, bold=True, color='FF1F4E78') if kind == 'title'
                  else Font(name='Arial', size=10, bold=True, color='FF1F4E78') if kind == 'h'
                  else Font(name='Consolas' if '  ' in t else 'Arial', size=9))
    ws.column_dimensions['A'].width = 110
    ws.sheet_view.showGridLines = False


peers_quarter_sheet()
hban_history_sheet()
camels_sheet()
estimates_sheet()
timeseries_sheet()
write_readme()
for ws in wb.worksheets[1:]:
    ws.sheet_view.zoomScale = 90
wb.save(OUT)
print('wrote', OUT)

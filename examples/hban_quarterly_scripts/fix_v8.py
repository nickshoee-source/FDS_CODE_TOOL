"""
V8 clean-up pass on top of V7.

 1. Remove ACL Ratio from Credit and Credit - Estimates (user request) --
    values AND formatting, so no empty formatted block is left behind.
 2. Delete ghost formatting: empty-but-formatted blocks left from earlier
    removals (Capital's old Dividends Paid block, Balance Sheet's old
    Average Deposits tail).
 3. Move loose test/scratch cells (outside the metric blocks) onto a new
    'Scratch Notes' sheet, stored verbatim as TEXT with their original
    address. Several now sat under the wrong metric after the V5 reorder.
 4. Valuation Market Cap FITB 2026/3F (AE27): a one-off FE_ESTIMATE(MV) test
    pointing at D8 while the other 10 banks use FF_MKT_VAL -> restored to
    the column's formula; the test is preserved on Scratch Notes.
 5. Error-guard every cell-reference calculation: =IFERROR(<math>,"-").
    FactSet fills missing quarters with "-" (text), so without the guard the
    Estimates sheets show #VALUE!, and divisions show #DIV/0!.
 6. TSR's final quarter referenced the blank anchor column for its end date.
    It now computes the next quarter from its own period label (verified in
    a real spreadsheet engine incl. the 4F -> next-year 1F rollover).
 7. Balance Sheet 'Earning Assets EOP' passed a blank row-6 cell as a 3rd
    argument. Dropped it -- an empty trailing FQL argument is the default, so
    results are unchanged, but a stray entry in row 6 can no longer silently
    alter all 308 formulas.
 8. Freeze panes normalised to E8 (keeps names/tickers + headers visible).
 9. Column grouping reset outside metric blocks; title bars fitted to content.

Deliberately NOT touched: Valuation's 2026/3F-4F FDSRC time-series spills
(user-built; the "numbers" in 2026/4F are their spill output, not pastes).
"""
import re
import copy
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

SRC = '/home/claude/Quarterly_Template_V7.xlsx'
OUT = '/home/claude/Quarterly_Template_V8.xlsx'
ORIG = '/home/claude/Quarterly_Template_V3_8-31-26_Edits.xlsx'   # user's own layout, for context

wb = load_workbook(SRC, data_only=False)
orig = load_workbook(ORIG, data_only=False)


def txt(c):
    v = c.value
    return v.text if isinstance(v, ArrayFormula) else v


def s_(v):
    return v if isinstance(v, str) else ''


def C(i):
    return get_column_letter(i)


def runs(ws, mr):
    out, cur, st = [], None, None
    for c in range(5, ws.max_column + 2):
        m = txt(ws.cell(row=mr, column=c))
        if not isinstance(m, str):
            m = None
        if m != cur:
            if cur is not None:
                out.append((cur, st, c - 1))
            cur, st = m, c
    return out


def wipe(ws, r, c):
    """Remove a cell completely (value + style)."""
    if (r, c) in ws._cells:
        del ws._cells[(r, c)]


def blocks_of(ws):
    return [cr for cr in (3, 22) if txt(ws.cell(row=cr + 1, column=5)) is not None]


LOG = []

# ------------------------------------------------------------ 1. ACL Ratio
for sheet in ['Credit', 'Credit - Estimates']:
    ws = wb[sheet]
    for lab, s, e in runs(ws, 4):
        if lab == 'ACL Ratio':
            for c in range(s, e + 1):
                for r in list(range(3, 6)) + list(range(8, 19)):
                    wipe(ws, r, c)
            LOG.append(f'{sheet}: removed ACL Ratio ({C(s)}:{C(e)})')

# ---------------------------------------------- 3/4. scratch -> Scratch Notes
scratch = []           # (sheet, addr, text, context)
for sheet in wb.sheetnames:
    ws = wb[sheet]
    bl = blocks_of(ws)
    legit = {}
    for cr in bl:
        cols = set()
        for lab, s, e in runs(ws, cr + 1):
            cols.update(range(s, e + 1))
        legit[cr] = cols
    for (r, c) in sorted(list(ws._cells.keys())):
        cell = ws._cells.get((r, c))
        if cell is None:
            continue
        v = txt(cell)
        if v is None:
            continue
        if r == 1:
            continue
        if sheet.startswith('Valuation') and (r, c) == (39, 3):
            continue                                        # share-price footnote
        ok = False
        for cr in bl:
            if r in {cr, cr + 1, cr + 2} | set(range(cr + 5, cr + 16)) and (c in legit[cr] or c in (3, 4)):
                ok = True
            if r == cr + 4 and c in (3, 4):
                ok = True
        if ok:
            continue
        # context: which metric that column sat under in the user's own 8/31 file
        ctx = ''
        if sheet in orig.sheetnames:
            ow = orig[sheet]
            mr = 23 if (r >= 22 and txt(ow.cell(row=23, column=5))) else 4
            m = txt(ow.cell(row=mr, column=c))
            ctx = f'under "{m}" in your 8/31 file' if isinstance(m, str) else ''
        arr = cell.value.ref if isinstance(cell.value, ArrayFormula) else ''
        scratch.append((sheet, f'{C(c)}{r}', str(v), ctx, arr))

# spilled cached values belonging to a scratch array anchor are recorded
# with the anchor, not as separate notes
anchor_ranges = {}
for sheet, addr, v, ctx, arr in scratch:
    if arr and ':' in arr:
        a, b = arr.split(':')
        anchor_ranges.setdefault(sheet, []).append((a, b))


def in_spill(sheet, addr):
    m = re.match(r'([A-Z]+)(\d+)', addr)
    col, row = column_index_from_string(m.group(1)), int(m.group(2))
    for a, b in anchor_ranges.get(sheet, []):
        ma, mb = re.match(r'([A-Z]+)(\d+)', a), re.match(r'([A-Z]+)(\d+)', b)
        c1, r1 = column_index_from_string(ma.group(1)), int(ma.group(2))
        c2, r2 = column_index_from_string(mb.group(1)), int(mb.group(2))
        if r1 <= row <= r2 and c1 <= col <= c2 and addr != a:
            return True
    return False


notes = [x for x in scratch if not in_spill(x[0], x[1])]

# Market Cap AE27 test cell (inside the data block)
val = wb['Valuation']
ae27 = txt(val['AE27'])
if 'FE_ESTIMATE(MV' in s_(ae27):
    notes.append(('Valuation', 'AE27', ae27,
                  'Market Cap, FITB, 2026/3F -- the only bank in that column not on FF_MKT_VAL', ''))
    val['AE27'] = '=FDS(D27,"FF_MKT_VAL(QTR,"&AE$5&")")'
    val['AE27'].number_format = '0.00'
    LOG.append('Valuation: AE27 restored to FF_MKT_VAL like the rest of its column')

# now remove all scratch cells (and their spill targets) from their sheets
for sheet, addr, v, ctx, arr in scratch:
    ws = wb[sheet]
    m = re.match(r'([A-Z]+)(\d+)', addr)
    wipe(ws, int(m.group(2)), column_index_from_string(m.group(1)))

# build the Scratch Notes sheet
sn = wb.create_sheet('Scratch Notes')
sn.sheet_view.showGridLines = False
hdr_font = Font(name='Arial', bold=True, size=10, color='FFFFFF')
hdr_fill = PatternFill('solid', fgColor='1F4E78')
thin = Side(style='thin', color='D9D9D9')
brd = Border(left=thin, right=thin, top=thin, bottom=thin)
sn['A1'] = 'Scratch Notes — test cells moved out of the data sheets (stored as text, not live)'
sn['A1'].font = Font(name='Arial', bold=True, size=13, color='FFFFFF')
sn['A1'].fill = hdr_fill
sn.merge_cells('A1:E1')
sn.row_dimensions[1].height = 22
sn['A2'] = ('These sat outside the metric blocks (or, for Valuation AE27, were a one-off test '
            'inside one). Several ended up under a different metric after the columns were '
            're-ordered, so they were moved here verbatim. Copy any back as needed.')
sn['A2'].font = Font(name='Arial', italic=True, size=9, color='595959')
sn['A2'].alignment = Alignment(wrap_text=True, vertical='top')
sn.merge_cells('A2:E2')
sn.row_dimensions[2].height = 30
for i, h in enumerate(['Sheet', 'Cell', 'Context', 'Array range', 'Content (verbatim)'], 1):
    c = sn.cell(row=4, column=i, value=h)
    c.font, c.fill, c.border = hdr_font, hdr_fill, brd
    c.alignment = Alignment(horizontal='center', vertical='center')
for i, (sheet, addr, v, ctx, arr) in enumerate(notes, start=5):
    vals = [sheet, addr, ctx, arr if (arr and ':' in arr) else '', v.replace('_xll.', '')]
    for j, x in enumerate(vals, 1):
        c = sn.cell(row=i, column=j, value=x)
        c.data_type = 's'                                   # text, never evaluated
        c.font = Font(name='Consolas' if j == 5 else 'Arial', size=9)
        c.border = brd
        c.alignment = Alignment(vertical='top', wrap_text=(j in (3, 5)))
for col, w in zip('ABCDE', (22, 7, 38, 12, 110)):
    sn.column_dimensions[col].width = w
sn.freeze_panes = 'A5'
LOG.append(f'Scratch Notes: {len(notes)} test cells preserved; '
           f'{len(scratch) - len(notes)} cached spill values discarded with their anchors')

# ------------------------------------------------ 5. IFERROR on cell-math
CELLMATH = {('Shareholder Returns', 'Buyback Yield'), ('Shareholder Returns', 'Capital Return Yield'),
            ('Revenue & Fee Mix', 'Total Revenue'), ('NIM & Margin', 'Cost of Funds'),
            ('Balance Sheet Evolution', 'TBV')}
EST_BASE = {'Balance Sheet Evol': 'Balance Sheet Evolution'}
wrapped = 0
for sheet in wb.sheetnames:
    ws = wb[sheet]
    base = sheet.replace(' - Estimates', '')
    base = EST_BASE.get(base, base)
    for cr in blocks_of(ws):
        for lab, s, e in runs(ws, cr + 1):
            if (base, lab) not in CELLMATH:
                continue
            for col in range(s, e):                          # skip anchor col
                for r in range(cr + 5, cr + 16):
                    cell = ws.cell(row=r, column=col)
                    f = s_(txt(cell))
                    if f.startswith('=') and not f.upper().startswith('=IFERROR('):
                        cell.value = f'=IFERROR({f[1:]},"-")'
                        wrapped += 1
LOG.append(f'Error-guarded {wrapped} cell-math formulas with IFERROR(...,"-")')

# ------------------------------------------------ 6. TSR final quarter
for sheet in ['Shareholder Returns', 'Shareholder Returns - Estimates']:
    ws = wb[sheet]
    for lab, s, e in runs(ws, 4):
        if lab != 'TSR':
            continue
        lastc = e - 1                                        # last period column
        L, A = C(lastc), C(e)                                # A = anchor column
        nxt = (f'(VALUE(LEFT({L}$5,4))+(MID({L}$5,6,1)="4"))&"/"&'
               f'(MOD(VALUE(MID({L}$5,6,1)),4)+1)&"F')
        for r in range(8, 19):
            cell = ws.cell(row=r, column=lastc)
            f = s_(txt(cell))
            new = f.replace(f'"&{A}$5&")/100")', f'"&{nxt})/100")')
            if new == f:
                raise SystemExit(f'TSR pattern not found in {sheet}!{L}{r}: {f}')
            cell.value = new.replace('_xll.', '')
        LOG.append(f'{sheet}: TSR {L} (last quarter) now computes its own end date')

# ------------------------------------------------ 7. Earning Assets EOP
ws = wb['Balance Sheet Evolution']
fixed = 0
for lab, s, e in runs(ws, 4):
    if lab != 'Earning Assets EOP':
        continue
    for col in range(s, e):
        for r in range(8, 19):
            cell = ws.cell(row=r, column=col)
            f = s_(txt(cell))
            new = re.sub(r',"&[A-Z]{1,3}\$6&"', '', f).replace('_xll.', '')
            if new != f:
                cell.value = new
                fixed += 1
LOG.append(f'Balance Sheet Evolution: dropped blank row-6 argument from {fixed} Earning Assets formulas')

# ----------------------------- 2/8/9. ghost formatting, grouping, freeze, title
for sheet in wb.sheetnames:
    if sheet == 'Scratch Notes':
        continue
    ws = wb[sheet]
    bl = blocks_of(ws)
    legit, all_metric_cols = {}, set()
    for cr in bl:
        cols = set()
        for lab, s, e in runs(ws, cr + 1):
            cols.update(range(s, e + 1))
        legit[cr] = cols
        all_metric_cols |= cols
    ghosts = 0
    for (r, c) in list(ws._cells.keys()):
        cell = ws._cells[(r, c)]
        if cell.value is not None or r == 1:
            continue
        keep = False
        for cr in bl:
            if r in {cr, cr + 1, cr + 2} | set(range(cr + 5, cr + 16)) and (c in legit[cr] or c in (3, 4)):
                keep = True
            if r == cr + 4 and c in (3, 4):
                keep = True
        if not keep:
            del ws._cells[(r, c)]
            ghosts += 1
    if ghosts:
        LOG.append(f'{sheet}: removed {ghosts} empty formatted (ghost) cells')
    # grouping outside metric blocks
    for L, dim in list(ws.column_dimensions.items()):
        try:
            ci = column_index_from_string(L)
        except ValueError:
            continue
        if ci >= 5 and ci not in all_metric_cols:
            dim.outlineLevel = 0
            dim.hidden = False
    # freeze
    if ws.freeze_panes != 'E8':
        LOG.append(f'{sheet}: freeze pane {ws.freeze_panes} -> E8')
        ws.freeze_panes = 'E8'
    # title bar
    last = max(all_metric_cols) if all_metric_cols else 4
    style = copy.copy(ws.cell(row=1, column=1)._style)
    for m in [m for m in ws.merged_cells.ranges if m.min_row == 1]:
        ws.unmerge_cells(str(m))
    for c in range(1, last + 1):
        ws.cell(row=1, column=c)._style = copy.copy(style)
    for c in range(last + 1, ws.max_column + 1):
        wipe(ws, 1, c)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)

wb.save(OUT)
print('saved', OUT, '\n')
for l in LOG:
    print(' -', l)
print('\nScratch Notes contents:')
for sheet, addr, v, ctx, arr in notes:
    print(f'   {sheet}!{addr:<5} {ctx[:48]:<48} {v[:60]}')

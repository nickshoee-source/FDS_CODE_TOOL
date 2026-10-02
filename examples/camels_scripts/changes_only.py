import zipfile, shutil
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.worksheet.formula import ArrayFormula
SRC='/mnt/user-data/outputs/CAMELS_FDS_Code_Analysis_V5.xlsx'
ORIG='/home/claude/camels/original.xlsx'
OUT='/mnt/user-data/outputs/CAMELS_FDS_Code_Analysis_V5_Changes_Highlighted.xlsx'
wb=load_workbook(SRC); o=load_workbook(ORIG)
YEL=PatternFill('solid',fgColor='FFFFFF00'); NOFILL=PatternFill(fill_type=None)
def norm(v):
    v=v.text if isinstance(v,ArrayFormula) else v
    if isinstance(v,str): v=v.replace('_xll.','')
    return v
isy=lambda c: c.fill.fill_type=='solid' and c.fill.fgColor.type=='rgb' and c.fill.fgColor.rgb=='FFFFFF00'
stats={}
for name in o.sheetnames:
    ws,ow=wb[name],o[name]; add=rem=chg=0
    for r in range(1,max(ws.max_row,ow.max_row)+1):
        for c in range(1,max(ws.max_column,ow.max_column)+1):
            cell=ws.cell(r,c); changed = norm(cell.value)!=norm(ow.cell(r,c).value)
            chg+=changed
            if changed and not isy(cell): cell.fill=YEL; add+=1
            elif not changed and isy(cell): cell.fill=NOFILL; rem+=1
    stats[name]=(chg,add,rem)
FM=wb['Formatted Annual Camels']; FM['A1'].value='$ in millions (USD). Yellow = changed from the original file.'
# change log is a new tab: keep yellow meaning only "changed cell" -> flag confirms in bold red text instead
for row in wb['Change Log'].iter_rows(min_row=2):
    for c in row:
        if isy(c): c.fill=NOFILL; c.font=Font(name='Arial',size=9,bold=True,color='FFB00B1C')
wb.save(OUT)
tmp=OUT+'.tmp'
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
    for it in zin.infolist():
        d=zin.read(it.filename)
        if it.filename=='[Content_Types].xml' and b'LabelInfo' not in d:
            d=d.replace(b'</Types>',b'<Override PartName="/docMetadata/LabelInfo.xml" ContentType="application/vnd.ms-office.classificationlabels+xml"/></Types>')
        if it.filename=='_rels/.rels' and b'LabelInfo' not in d:
            d=d.replace(b'</Relationships>',b'<Relationship Type="http://schemas.microsoft.com/office/2020/02/relationships/classificationlabels" Target="docMetadata/LabelInfo.xml" Id="rIdLbl1"/></Relationships>')
        if it.filename!='docMetadata/LabelInfo.xml': zout.writestr(it,d)
    zout.writestr('docMetadata/LabelInfo.xml',zipfile.ZipFile(ORIG).read('docMetadata/LabelInfo.xml'))
shutil.move(tmp,OUT)
for k,v in stats.items(): print(f'{k:28} changed={v[0]:6}  yellow added={v[1]:4}  yellow removed (unchanged)={v[2]}')
# verify: yellow <=> changed
wb2=load_workbook(OUT)
for name in o.sheetnames:
    ws,ow=wb2[name],o[name]; bad=0
    for r in range(1,max(ws.max_row,ow.max_row)+1):
        for c in range(1,max(ws.max_column,ow.max_column)+1):
            cell=ws.cell(r,c)
            if (norm(cell.value)!=norm(ow.cell(r,c).value)) != isy(cell): bad+=1
    print(name,'mismatches:',bad)

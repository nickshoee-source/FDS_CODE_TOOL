"""V5: roll the FFI@FF@FB code chains out to all 669 banks and carry the
CAMELS HISTORICAL ANNUAL (red theme) colour scheme/layout through every sheet."""
import re, zipfile, shutil
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.formula import ArrayFormula

SRC = '/mnt/user-data/outputs/CAMELS_FDS_Code_Analysis_V3.xlsx'
ORIG = '/home/claude/camels/original.xlsx'
OUT = '/mnt/user-data/outputs/CAMELS_FDS_Code_Analysis_V5.xlsx'
wb = load_workbook(SRC)
orig = load_workbook(ORIG)
H, FM = wb['CAMELS HISTORICAL ANNUAL'], wb['Formatted Annual Camels']
FIRST, LAST = 3, 671

# ------------------------------------------------ colours (same as historical sheet)
RED = 'FFB00B1C'
RED_FILL = PatternFill('solid', fgColor=RED)
TINT = PatternFill('solid', fgColor='FFF7E1E4')
NOFILL = PatternFill(fill_type=None)
YEL = PatternFill('solid', fgColor='FFFFFF00')
WHITE_B9 = Font(name='Arial', size=9, bold=True, color='FFFFFFFF')
RED_B9 = Font(name='Arial', size=9, bold=True, color=RED)
RED_B10 = Font(name='Arial', size=10, bold=True, color=RED)
BODY = Font(name='Arial', size=10, color='FF262626')
LINE = Border(bottom=Side(style='medium', color=RED))
CENTER = Alignment(horizontal='center', vertical='center')

# ------------------------------------------------ code chains
I = lambda c, u=False: f'{c}(ANN_L,{{Y}}' + (',,,,USD)' if u else ')')      # FactSet Fundamentals Industry
F = lambda c, u=False: f'{c}(ANN_R,{{Y}}' + (',,,,USD)' if u else ')')      # FactSet Fundamentals
B = lambda c, u=False: f'{c}(ANN,{{Y}},,,RF' + (',USD)' if u else ')')      # FactSet regulatory (US banks)
ch = lambda *x: '@'.join(x)
LOANS = ch(I('FFI_LOAN_ADV_TOT', 1), F('FF_BK_LOAN_TOT', 1), B('FB_TOT_HFI_HFS_UNEARN_INC', 1))
DEPS = ch(I('FFI_DEPS_TOT', 1), F('FF_DEPS', 1), B('FB_DEPS', 1))
CODES = {
    6: ch(I('FFI_COM_EQ_TIER1_TOT', 1), F('FF_BK_COM_EQ_TIER1_TOT', 1), B('FB_COM_EQ_TIER1', 1)),
    7: ch(I('FFI_COM_EQ_TIER1_RATIO'), F('FF_BK_COM_EQ_TIER1_RATIO'), B('FB_COM_EQ_TIER1_RATIO')),
    8: ch(I('FFI_TIER1_CAP', 1), F('FF_TIER1_CAP', 1), B('FB_TIER1_CAP', 1)),
    9: ch(I('FFI_CAP_RATIO_TIER1'), F('FF_CAP_RATIO_TIER1'), B('FB_CAP_RATIO_TIER1')),
    10: ch(I('FFI_CAP_RATIO_TOT'), F('FF_CAP_RATIO_TOT'), B('FB_CAP_RATIO_TOT')),
    11: ch(I('FFI_LEV_RATIO_RPT'), F('FF_BK_LEV_RATIO'), B('FB_LEV_RATIO')),
    12: ch(I('FFI_RWA', 1), F('FF_ASSETS_RISK_WGHT', 1), B('FB_ASSETS_RISK_WGHT', 1)),
    14: ch(I('FFI_NPL_LOAN_ADV', 1), F('FF_LOAN_NONPERF', 1), B('FB_NONPERF_LOAN', 1)),
    15: ch(I('FFI_LOAN_LOSS_PROV', 1), F('FF_LOAN_LOSS_PROV', 1), B('FB_LOAN_LOSS_PROV', 1)),
    16: '(' + ch(I('FFI_NPL_LOAN_LOSS_RSRV_RATIO'), F('FF_NONPERF_LOAN_LOSS_RSRV')) + ')/100',
    17: '(' + ch(I('FFI_NPL_LOAN_RATIO'), F('FF_NONPERF_LOAN_PCT'), B('FB_NONPERF_LOAN_PCT')) + ')/100',
    19: ch(I('FFI_EFF_RATIO', 1), F('FF_BK_EFF_RATIO'), B('FB_EFF_RATIO')),
    21: ch(I('FFI_ROTE'), B('FB_ROTE')),
    22: ch(I('FFI_ROTCE'), F('FF_ROTCE'), B('FB_ROTCE')),
    23: ch(I('FFI_ROA'), F('FF_ROA'), B('FB_ROA')),
    24: ch(I('FFI_AVG_BAL_INT_RATE_NET_MGN'), F('FF_INT_MGN'), B('FB_INT_MGN')),
    25: ch(I('FFI_INT_INC_NET', 1), F('FF_INT_INC_NET', 1), B('FB_INT_INC_NET', 1)),
    26: ch(I('FFI_NON_INT_INC', 1), F('FF_NON_INT_INC', 1), B('FB_NON_INT_INC', 1)),
    28: LOANS, 29: DEPS, 30: f'({LOANS})/({DEPS})',
    31: ch(I('FFI_BK_LIQ_COVG_RATIO'), F('FF_BK_LIQ_COVG_RATIO')),
    32: I('FFI_BK_NSFR'),
    34: ch(I('FFI_LOAN_AMORT_CUST', 1), F('FF_LOAN_NET', 1), B('FB_TOT_HFI_HFS_UNEARN_INC', 1)),
    35: I('FFI_LOAN_AMORT_BK', 1),
    36: '(' + I('FFI_SECS_INVEST', 1) + '+' + I('FFI_TRADE_ACCT', 1) + '+' + I('FFI_DERIV_HEDGE', 1) + ')@' + F('FF_INVEST_TOT', 1),
    37: '(' + ch(I('FFI_DEPS', 1), F('FF_DEPS_CUST', 1)) + ')/(' + ch(I('FFI_DEPS_TOT', 1), F('FF_DEPS', 1), B('FB_DEPS', 1)) + ')',
}
CONFIRM = {6: 'FF_BK_COM_EQ_TIER1_TOT', 34: 'FF_LOAN_NET / FB_TOT_HFI_HFS_UNEARN_INC are stand-ins (net / total loans)',
           36: 'FF_INVEST_TOT; FFI sum is NA if any piece is missing', 37: 'FF_DEPS_CUST'}

for r, code in CODES.items():
    for c in range(FIRST, LAST + 1):
        col = L(c)
        H.cell(r, c).value = f'=FDSC("-",{col}2,"' + code.replace('{Y}', f'"&{col}4&"') + '")'
        H.cell(r, c).fill = YEL
    assert '{Y}' not in H.cell(r, FIRST).value

# ------------------------------------------------ change log: replace historical rows
CL = wb['Change Log']
keep = [list(r) for r in CL.iter_rows(min_row=2, values_only=True) if r[0] != 'CAMELS HISTORICAL ANNUAL' or r[1] == 38]
t = lambda v: (v.text if isinstance(v, ArrayFormula) else v)
new = []
for r, code in CODES.items():
    before = t(orig['CAMELS HISTORICAL ANNUAL'].cell(r, 3).value)
    before = before.replace('_xll.', '') if isinstance(before, str) else '(blank)'
    n = code.count('@') + 1 if r not in (30, 37) else None
    status = ('PLEASE CONFIRM: ' + CONFIRM[r]) if r in CONFIRM else 'Codes from your FFI / FF / FB reference files + HBAN work'
    new.append(['CAMELS HISTORICAL ANNUAL', r, H.cell(r, 2).value, f'C{r}:YU{r}', before, H.cell(r, FIRST).value,
                'FFI first, then FF_, then FB_ (regulatory, US banks) via @ fallback', status])
CL.delete_rows(2, CL.max_row)
for row in new + keep:
    CL.append(row)
for rr in CL.iter_rows(min_row=2):
    for c in rr:
        c.font = Font(name='Arial', size=9)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.fill = NOFILL
    if 'CONFIRM' in str(rr[7].value):
        rr[7].fill = YEL
for c in CL[1]:
    c.fill, c.font = RED_FILL, WHITE_B9
    c.alignment = Alignment(vertical='center', wrap_text=True)
CL.row_dimensions[1].height = 24
CL.sheet_properties.tabColor = RED

# ------------------------------------------------ formatted sheet -> red theme
SECTIONS_F = (3, 11, 16, 18, 25, 31)
METRIC_F = [r for r in range(4, 37) if r not in SECTIONS_F]
for c in range(1, 673):
    for r in range(1, 37):
        cell = FM.cell(r, c)
        if r == 2 and c >= 2:
            cell.fill, cell.font = RED_FILL, WHITE_B9
        elif r in SECTIONS_F and c >= 2:
            cell.fill = TINT
            cell.font = RED_B10 if c == 3 else RED_B9
            if c >= 3 and r == 3:
                cell.border = Border(bottom=Side(style='thin', color=RED))
        elif r in METRIC_F and c >= 2:
            is_yellow = cell.fill.fill_type == 'solid' and cell.fill.fgColor.rgb == 'FFFFFF00'
            if c == 3:
                cell.fill = YEL                                  # every metric's underlying code changed
            elif not is_yellow:
                cell.fill = NOFILL                               # drop grey banding, as on the historical sheet
            cell.font = Font(name='Arial', size=10, bold=(c == 2), color='FF262626')
for r in range(3, 37):
    FM.cell(r, 1).fill = NOFILL                          # leftover grey banding in column A
FM['C2'].font = Font(name='Arial', size=11, bold=True, color='FFFFFFFF')
FM['A1'].font = Font(name='Arial', italic=True, size=9, color=RED)
FM.sheet_properties.tabColor = RED
H.sheet_properties.tabColor = RED

# ------------------------------------------------ single-company sheets -> same layout
SECTIONS_S = (5, 13, 18, 20, 27, 33)
for name in ('CAMELS NTM', 'CAMELS ESTIMATE ANNUAL', 'CAMELS LTM'):
    ws = wb[name]
    for r in (1, 2):
        for c in (1, 2, 3):
            cell = ws.cell(r, c)
            cell.fill, cell.font = RED_FILL, WHITE_B9
            cell.alignment = Alignment(horizontal='left' if c < 3 else 'center', vertical='center')
    for c in (1, 2, 3):
        y = ws.cell(4, c)
        y.fill, y.font, y.border, y.alignment = TINT, RED_B9, LINE, CENTER
        for r in SECTIONS_S:
            s = ws.cell(r, c)
            s.fill = TINT
            s.font = RED_B10
    for r in range(6, 39):
        if r in SECTIONS_S:
            continue
        ws.cell(r, 2).font = BODY
        d = ws.cell(r, 3)
        d.font = BODY
        d.alignment = CENTER
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 44
    ws.column_dimensions['C'].width = 18
    ws.freeze_panes = 'C5'
    ws.sheet_properties.tabColor = RED

wb.save(OUT)

# put back the sensitivity label part openpyxl drops
tmp = OUT + '.tmp'
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == '[Content_Types].xml' and b'LabelInfo' not in data:
            data = data.replace(b'</Types>', b'<Override PartName="/docMetadata/LabelInfo.xml" '
                                b'ContentType="application/vnd.ms-office.classificationlabels+xml"/></Types>')
        if item.filename == '_rels/.rels' and b'LabelInfo' not in data:
            data = data.replace(b'</Relationships>', b'<Relationship Type="http://schemas.microsoft.com/office/2020/02/'
                                b'relationships/classificationlabels" Target="docMetadata/LabelInfo.xml" Id="rIdLbl1"/>'
                                b'</Relationships>')
        if item.filename != 'docMetadata/LabelInfo.xml':
            zout.writestr(item, data)
    zout.writestr('docMetadata/LabelInfo.xml', zipfile.ZipFile(ORIG).read('docMetadata/LabelInfo.xml'))
shutil.move(tmp, OUT)
print('saved', OUT)

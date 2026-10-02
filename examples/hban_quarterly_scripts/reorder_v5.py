"""
V5: put every referenced (input) metric IN FRONT OF the metric that computes
from it, on each worksheet that contains cell-reference math.

Method: snapshot each affected block completely (labels, categories, period
labels, per-cell styles, column widths, and a parameterised copy of every
FDS formula), clear the block, then re-lay it out in the new order and
regenerate all formulas for the new positions.

Formula handling:
  * FDS metrics  -> template derived from the metric's own current formula, so
                    the user's hand-edited FactSet code is preserved verbatim.
                    Any reference into the metric's own 29-column span becomes
                    an offset token ({P+0}, {P+1}, ...) so multi-column
                    formulas such as TSR's start/end date pair keep working.
  * cell-math    -> rebuilt from a declarative spec once new positions are known.
  * identifier   -> normalised to $D{row} (relative row, absolute column). This
                    removes the lone $D$8 on TSR row 8, which pointed at itself
                    and so was harmless, but would have broken if copied down.

Sheets without any cell-math (Valuation, Profitability, Efficiency, Deposit
Competition, Credit, Capital, Estimates, Franchise) are NOT touched here, so
the static pasted values on Valuation's 2026/4F columns survive untouched.
"""
import re
import copy
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter, column_index_from_string

SRC = '/home/claude/Quarterly_Template_V4.xlsx'
OUT = '/home/claude/Quarterly_Template_V5.xlsx'

N_PERIODS, BLOCK_W, N_BANKS = 28, 29, 11
wb = load_workbook(SRC, data_only=False)


def txt(c):
    v = c.value
    return v.text if isinstance(v, ArrayFormula) else v


def s_(v):
    return v if isinstance(v, str) else ''


def C(i):
    return get_column_letter(i)


def block_rows(cat_row):
    first = cat_row + 5
    return cat_row, cat_row + 1, cat_row + 2, first, first + N_BANKS - 1


def metric_runs(ws, met_row):
    out, cur, st = [], None, None
    for c in range(5, ws.max_column + 2):
        m = txt(ws.cell(row=met_row, column=c))
        if not isinstance(m, str):
            m = None
        if m != cur:
            if cur is not None:
                out.append((cur, st, c - 1))
            cur, st = m, c
    return out


# ------------------------------------------------------------ desired order
ORDER = {
    ('Shareholder Returns', 3): [
        'TSR', 'Dividend Yield', 'Dividend Payout Ratio', 'Dividends Paid',
        # --- Buyback Yield inputs, then the calc
        'Share Repurchases', 'Market Cap (QTR_R)', 'Buyback Yield',
        # --- Capital Return Yield inputs, then the calc
        'Dividend Per Share', 'Common Shares Outstanding', 'Stock Repurchase CF',
        'Market Value (Company)', 'Capital Return Yield',
    ],
    ('Revenue & Fee Mix', 3): ['NII', 'Non-Interest Income', 'Total Revenue'],
    ('NIM & Margin', 3): [
        'NIM', 'EA Yield',
        'Interest Expense Total', 'Average Interest Bearing Liabilities',
        'Cost of Funds',
        'Deposit Cost', 'Interest Bearing Deposit Cost', 'Total Funding Cost',
        'Net Interest Spread', 'Taxable Equivalent NIM',
    ],
    # TBV's three inputs move to the front of the upper (Assets) block, which
    # sits above the Capital block where TBV lives -> inputs come first in
    # reading order.
    ('Balance Sheet Evolution', 3): [
        'Total Assets EOP', 'Intangible Assets', 'Total Liabilities',
        'Average Assets', 'Earning Assets EOP', 'Total Loans EOP',
    ],
    ('Balance Sheet Evolution', 22): [
        'Total Deposits EOP', 'Tangible Common Equity',
        'Average Tangible Common Equity', 'TBV', 'TBVPS',
    ],
}

# ----------------------------------------------- cell-math rebuild recipes
def cm_buyback(P, r, i):
    return f"={C(P[(3,'Share Repurchases')]+i)}{r}/{C(P[(3,'Market Cap (QTR_R)')]+i)}{r}"

def cm_capreturn(P, r, i):
    return (f"=(({C(P[(3,'Dividend Per Share')]+i)}{r}"
            f"*{C(P[(3,'Common Shares Outstanding')]+i)}{r})"
            f"+{C(P[(3,'Stock Repurchase CF')]+i)}{r})"
            f"/{C(P[(3,'Market Value (Company)')]+i)}{r}")

def cm_totalrev(P, r, i):
    return f"={C(P[(3,'NII')]+i)}{r}+{C(P[(3,'Non-Interest Income')]+i)}{r}"

def cm_cof(P, r, i):
    return (f"=(({C(P[(3,'Interest Expense Total')]+i)}{r}"
            f"/{C(P[(3,'Average Interest Bearing Liabilities')]+i)}{r})*100)*4")

def cm_tbv(P, r, i):
    rr = r - 19                      # Capital block row -> matching Assets block row
    return (f"={C(P[(3,'Total Assets EOP')]+i)}{rr}"
            f"-{C(P[(3,'Intangible Assets')]+i)}{rr}"
            f"-{C(P[(3,'Total Liabilities')]+i)}{rr}")

CELLMATH = {
    ('Shareholder Returns', 3, 'Buyback Yield'): cm_buyback,
    ('Shareholder Returns', 3, 'Capital Return Yield'): cm_capreturn,
    ('Revenue & Fee Mix', 3, 'Total Revenue'): cm_totalrev,
    ('NIM & Margin', 3, 'Cost of Funds'): cm_cof,
    ('Balance Sheet Evolution', 22, 'TBV'): cm_tbv,
}

ID_RE = re.compile(r'(?<![A-Za-z_])\$?D\$?\d+')
PER_RE = re.compile(r'(?<![A-Z$])(\$?)([A-Z]{1,3})(\$?)(\d+)')


def parameterize(formula, own_start):
    """FDS formula -> template with {ROW} and {P+k} tokens."""
    t = s_(formula).replace('_xll.', '')
    t = ID_RE.sub('$D{ROW}', t)

    def repl(m):
        d1, col, d2, row = m.groups()
        ci = column_index_from_string(col)
        k = ci - own_start
        if 0 <= k <= N_PERIODS:          # inside this metric's own span
            return f'{d1}{{P+{k}}}{d2}{row}'
        return m.group(0)
    return PER_RE.sub(repl, t)


def render(template, row, start_col, i):
    out = template.replace('{ROW}', str(row))
    for k in range(N_PERIODS + 1):
        tok = f'{{P+{k}}}'
        if tok in out:
            out = out.replace(tok, C(start_col + i + k))
    return out


CHANGELOG = []

for (sheet, cat_row) in [k for k in ORDER]:
    pass  # ordering handled below, grouped by sheet

# process sheet by sheet so block 3 positions exist before block 22 is built
by_sheet = {}
for (sheet, cat_row) in ORDER:
    by_sheet.setdefault(sheet, []).append(cat_row)

for sheet, cat_rows in by_sheet.items():
    ws = wb[sheet]
    positions = {}
    for cat_row in sorted(cat_rows):
        _, met_row, per_row, first, last = block_rows(cat_row)
        rows_all = [cat_row, met_row, per_row] + list(range(first, last + 1))

        # ---------------- snapshot -------------------------------------
        snap = {}
        for lab, s, e in metric_runs(ws, met_row):
            width = e - s + 1
            assert width == BLOCK_W, f"{sheet} {lab} width {width}"
            is_cm = (sheet, cat_row, lab) in CELLMATH
            snap[lab] = {
                'category': txt(ws.cell(row=cat_row, column=s)),
                'periods': [txt(ws.cell(row=per_row, column=s + i)) for i in range(BLOCK_W)],
                'template': None if is_cm else parameterize(
                    txt(ws.cell(row=first, column=s)), s),
                'is_cm': is_cm,
                'styles': [{r: copy.copy(ws.cell(row=r, column=s + i)._style)
                            for r in rows_all} for i in range(BLOCK_W)],
                'widths': [ws.column_dimensions[C(s + i)].width for i in range(BLOCK_W)],
                'numfmt': [{r: ws.cell(row=r, column=s + i).number_format
                            for r in range(first, last + 1)} for i in range(BLOCK_W)],
            }

        order = ORDER[(sheet, cat_row)]
        missing = set(snap) ^ set(order)
        assert not missing, f"{sheet} block {cat_row}: order mismatch {missing}"

        # ---------------- clear the block ------------------------------
        span_end = 4 + BLOCK_W * len(snap)
        for c in range(5, span_end + 1):
            for r in rows_all:
                ws.cell(row=r, column=c).value = None

        # ---------------- assign new positions -------------------------
        for idx, lab in enumerate(order):
            positions[(cat_row, lab)] = 5 + idx * BLOCK_W

        # ---------------- write back in the new order ------------------
        for lab in order:
            m = snap[lab]
            s = positions[(cat_row, lab)]
            for i in range(BLOCK_W):
                col = s + i
                for r, st in m['styles'][i].items():
                    ws.cell(row=r, column=col)._style = copy.copy(st)
                ws.cell(row=cat_row, column=col).value = m['category']
                ws.cell(row=met_row, column=col).value = lab
                ws.cell(row=per_row, column=col).value = m['periods'][i]
                if m['widths'][i]:
                    ws.column_dimensions[C(col)].width = m['widths'][i]
                if i < N_PERIODS:
                    for r in range(first, last + 1):
                        cell = ws.cell(row=r, column=col)
                        cell.value = (None if m['is_cm']
                                      else render(m['template'], r, s, i))
                        cell.number_format = m['numfmt'][i].get(r, '0.00')

        # ---------------- now fill the cell-math metrics ---------------
        for lab in order:
            if not snap[lab]['is_cm']:
                continue
            fn = CELLMATH[(sheet, cat_row, lab)]
            s = positions[(cat_row, lab)]
            for i in range(N_PERIODS):
                for r in range(first, last + 1):
                    cell = ws.cell(row=r, column=s + i)
                    cell.value = fn(positions, r, i)
                    cell.number_format = '0.00'

        CHANGELOG.append(
            (sheet, cat_row, [(lab, C(positions[(cat_row, lab)])) for lab in order]))

# --------------------------------------------------- restore column grouping
for name in wb.sheetnames:
    ws = wb[name]
    for met_row in (4, 23):
        if txt(ws.cell(row=met_row, column=5)) is None:
            continue
        for lab, s, e in metric_runs(ws, met_row):
            for i in range(e - s + 1):
                d = ws.column_dimensions[C(s + i)]
                d.outlineLevel = 1 if i < N_PERIODS else 0
                d.hidden = False
    if ws.sheet_properties.outlinePr is not None:
        ws.sheet_properties.outlinePr.summaryRight = True

wb.save(OUT)
print('saved', OUT, '\n')
for sheet, cat_row, order in CHANGELOG:
    print(f'{sheet}  (block starting row {cat_row})')
    for lab, col in order:
        print(f'    {col:>3}  {lab}')
    print()

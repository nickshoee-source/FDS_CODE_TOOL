"""
Add a 'Data Check' sheet to V10: every value from the user's 8/31 FactSet
refresh that is probably wrong, needs a look, or is unusual-but-real, each
with its worksheet, bank, quarter, a clickable link to the V10 cell, and the
value as this file displays it (e.g. NCO x4).
"""
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.hyperlink import Hyperlink

SRC = '/home/claude/Quarterly_Template_V3_8-31-26_Edits.xlsx'
OUT = '/home/claude/Quarterly_Template_V10.xlsx'
V = load_workbook(SRC, data_only=True)
F = load_workbook(SRC)
wb = load_workbook(OUT)

PER = [f'{y}/{q}F' for y in range(2020, 2027) for q in range(1, 5)]
ROW = {'FITB': 0, 'ZION': 1, 'RF': 2, 'PNC': 3, 'MTB': 4, 'FHN': 5, 'HBAN': 6,
       'USB': 7, 'KEY': 8, 'CFG': 9, 'TFC': 10}
SCALE = {'Efficiency Ratio': 100, 'TSR': 100, 'Buyback Yield': 100,
         'Capital Return Yield': 100, 'NCO Ratio': 4, 'Deposit Cost': 4}
KIND = {'Total Loans EOP': '$', 'Efficiency Ratio': '%', 'NCO Ratio': '%', 'Deposit Cost': '%',
        'Cost of Funds': '%', 'Interest Bearing Deposit Cost': '%', 'Total Funding Cost': '%',
        'Taxable Equivalent NIM': '%', 'Non-Interest Income': '$', 'Capital Return Yield': '%',
        'Total Deposits': '$', 'Enterprise Value': '$', 'Net Income': '$', 'P/E': 'x',
        'Dividend Payout Ratio': '%', 'Total Assets EOP': '$', 'Branch Count': '#'}


def t(c):
    v = c.value
    return v.text if isinstance(v, ArrayFormula) else v


def runs(ws, mr):
    out, cur, s0 = [], None, None
    for c in range(5, ws.max_column + 2):
        m = t(ws.cell(row=mr, column=c))
        if not isinstance(m, str):
            m = None
        if m != cur:
            if cur is not None:
                out.append((cur, s0, c - 1))
            cur, s0 = m, c
    return out


def block_of(ws, metric):
    for cr in (3, 22):
        for lab, s0, e in runs(ws, cr + 1):
            if lab == metric:
                return cr, s0
    raise KeyError(metric)


def old_value(sheet, metric, bank, period):
    cr, s0 = block_of(F[sheet], metric)
    v = V[sheet].cell(row=cr + 5 + ROW[bank], column=s0 + PER.index(period)).value
    return v * SCALE.get(metric, 1) if isinstance(v, (int, float)) else v


def new_cell(sheet, metric, bank=None, period=None):
    cr, s0 = block_of(wb[sheet], metric)
    r = cr + 5 + (ROW[bank] if bank else 0)
    c = s0 + (PER.index(period) if period else 0)
    return f'{get_column_letter(c)}{r}'


def fmt(metric, v):
    if not isinstance(v, (int, float)):
        return str(v)
    k = KIND.get(metric, '')
    neg = v < 0
    a = abs(v)
    if k == 'x':
        return 'NM' if (v > 100 or v < 0) else f'{v:,.2f}x'
    s = {'$': f'${a:,.2f}', '%': f'{a:,.2f}%', '#': f'{a:,.0f}'}.get(k, f'{a:,.2f}')
    return f'({s})' if neg else s


def vals(sheet, metric, bank, periods):
    return ' / '.join(fmt(metric, old_value(sheet, metric, bank, p)) for p in periods)


# ------------------------------------------------------------------ findings
PROB, CHECK, REAL, FIXED = 'Probably incorrect', 'Check', 'Unusual but real', 'Already fixed'
rows = []


def add(status, sheet, metric, bank='', periods=(), why='', cause='', link=None, value=None):
    periods = list(periods)
    cell = link or (new_cell(sheet, metric, bank or None, periods[0] if periods else None))
    shown = value if value is not None else (vals(sheet, metric, bank, periods) if bank and periods else '')
    rows.append((status, sheet, metric, bank, ', '.join(periods) if len(periods) < 4
                 else f'{periods[0]} to {periods[-1]}', cell, shown, why, cause))


NIM = 'NIM & Margin'
# ---- probably incorrect
add(PROB, 'Balance Sheet Evolution', 'Total Loans EOP', 'FITB', ['2026/1F', '2026/2F'],
    'Loans are larger than total assets ($297,039 / $300,180) - impossible',
    'Bad FactSet figure for FITB after the Comerica deal (expect ~$175B). Check FF_BK_LOAN_TOT for FITB or use another loans code.')
add(PROB, 'Efficiency', 'Efficiency Ratio', 'USB', ['2020/2F'],
    '624% while the quarters around it are 57-59%',
    'Decimal slip in FactSet data (6.24 instead of ~0.60).')
add(PROB, 'Credit', 'NCO Ratio', 'USB', ['2020/2F'],
    'Negative (-0.55%) between +0.53% and +0.66% - the sign has flipped',
    'USB had net charge-offs, not recoveries, in Q2 2020. FactSet sign error.')
add(PROB, NIM, 'Deposit Cost', 'PNC', ['2025/4F'],
    '1,696% while the quarters around it are ~1.5-1.8%',
    'Bad FactSet data point.')
add(PROB, NIM, 'Cost of Funds', 'RF', ['2026/1F'],
    '33.13% while the quarters around it are ~2%',
    'One of its two inputs is wrong that quarter - check Interest Expense Total and Avg Interest Bearing Liabilities.')
for m in ('Cost of Funds', 'Interest Bearing Deposit Cost', 'Total Funding Cost'):
    add(PROB, NIM, m, 'HBAN', ['2021/1F'],
        'Negative funding cost (-0.53%) is impossible; NIM is also above the asset yield that quarter',
        'Bad FactSet data for HBAN Q1 2021 margin items.')
add(PROB, NIM, 'Taxable Equivalent NIM', 'CFG', ['2026/1F', '2026/2F'],
    'Exactly double CFG\'s plain NIM (3.06% / 3.13%); earlier quarters sit a few bps above NIM',
    'Bad FactSet data.')
add(PROB, NIM, 'Interest Bearing Deposit Cost', why='Equals Cost of Funds in all 286 cells (median gap 0.008 pts)',
    cause='FF_BK_INT_COST_INTB_AVG measures ALL interest-bearing liabilities, not deposits. Needs a deposit-only code.',
    value='whole column')
add(PROB, NIM, 'Total Funding Cost', why='Identical to Interest Bearing Deposit Cost in every cell',
    cause='FF_BK_COST(QTR_NINTB_LIABS) returns the same series. Needs a different code.', value='whole column')

# ---- check
add(CHECK, 'Efficiency', 'Efficiency Ratio', 'MTB', ['2026/2F'], 'Jumps to 105.6% from ~54-58%',
    'Possibly a one-off charge - confirm against MTB\'s Q2 2026 results.')
add(CHECK, 'Efficiency', 'Efficiency Ratio', 'CFG', ['2024/4F'], 'Jumps to 101.3% from ~66-68%',
    'Confirm against CFG\'s Q4 2024 results.')
add(CHECK, 'Revenue & Fee Mix', 'Non-Interest Income', 'ZION', ['2026/2F'], 'Jumps to $450 from ~$175-200',
    'Possibly a one-time gain - confirm.')
add(CHECK, 'Shareholder Returns', 'Capital Return Yield', 'HBAN', ['2021/3F'],
    '5.89% although buybacks alone were 2.24% that quarter',
    'Stock Repurchase CF may include preferred stock redemptions.')
for bank, ps in [('ZION', ['2020/1F', '2020/2F', '2020/3F']), ('ZION', ['2023/1F', '2023/2F', '2023/3F', '2023/4F', '2024/1F', '2024/2F']),
                 ('KEY', ['2020/1F', '2020/2F', '2020/3F']), ('PNC', ['2022/1F'])]:
    add(CHECK, 'Deposit Competition', 'Total Deposits', bank, ps,
        'NIBD + interest-bearing deposits come to 5-16% less than total deposits',
        'FactSet\'s split leaves out some deposit types (e.g. brokered/time).', value='')
add(CHECK, 'Valuation', 'Enterprise Value', 'RF',
    ['2021/1F', '2021/2F', '2021/3F', '2021/4F', '2022/1F'], 'Negative for five quarters',
    'Cash exceeded market cap + debt. EV is not a meaningful measure for banks.')
for sheet, metric, why in [
        ('Balance Sheet Evolution', 'Total Loans EOP', 'RF is missing 25 of 28 quarters'),
        ('Estimates', 'Consensus CET1', 'MTB is missing 27 of 28 quarters'),
        ('Balance Sheet Evolution', 'Average Tangible Common Equity', 'PNC is missing 18 of 28 quarters'),
        ('Shareholder Returns', 'Buyback Yield', 'ZION missing 13, HBAN 10 quarters - likely quarters with no buybacks, so "-" shows instead of 0.00%'),
        ('Franchise', 'Branch Count', 'KEY missing 4 quarters')]:
    add(CHECK, sheet, metric, why=f'Data gap: {why}', cause='FactSet has no value for these - shows "-".', value='gaps')

# ---- unusual but real
for sheet, metric, bank, ps, why in [
        ('Valuation', 'P/E', 'KEY', ['2024/3F'], 'KeyCorp securities-repositioning losses (Q3 and Q4 2024): net losses, efficiency ~140-156%, P/E 1,523x (now shows NM), payout 7,455%'),
        ('Profitability', 'Net Income', 'TFC', ['2023/4F'], 'Truist $6.1B goodwill impairment: net loss -$5,090, ROE -32%, efficiency 196%'),
        ('Revenue & Fee Mix', 'Non-Interest Income', 'TFC', ['2024/2F'], 'Truist securities-repositioning loss: fee income -$5,192, revenue -$1,665, net income -$3,925'),
        ('Profitability', 'Net Income', 'PNC', ['2020/2F'], 'COVID reserve builds: PNC and RF net losses in Q2 2020'),
        ('Profitability', 'Net Income', 'HBAN', ['2021/2F'], 'TCF merger (June 2021): -$15 net income, balance sheet +39%'),
        ('Balance Sheet Evolution', 'Total Assets EOP', 'FHN', ['2020/3F'], 'IBERIABANK merger: balance sheet +71%, bargain-purchase gain lifts fee income and ROE'),
        ('Balance Sheet Evolution', 'Total Assets EOP', 'MTB', ['2022/2F'], 'People\'s United acquisition: balance sheet +36%'),
        ('Balance Sheet Evolution', 'Total Assets EOP', 'FITB', ['2026/1F'], 'Comerica acquisition: assets +39%, deposits +36%, RWA +46%, branches +36%'),
        ('Franchise', 'Branch Count', 'HBAN', ['2026/1F'], 'Cadence acquisition: branches +39%'),
        ('Credit', 'NCO Ratio', 'ZION', ['2021/2F'], 'Small negative NCOs for ZION and FHN in 2021 are net recoveries')]:
    add(REAL, sheet, metric, bank, ps, why, 'Real event - not a data error.')

# ---- fixed in this file
for sheet, metric, why in [
        ('Credit', 'NCO Ratio', 'Was quarterly (~4x too low) - now annualized (x4). The Estimates-sheet NCO was already annualized, so it was left alone.'),
        (NIM, 'Deposit Cost', 'Was quarterly - now annualized (x4), including the reported-value branch on the Estimates sheet'),
        ('Profitability', 'Effective Tax Rate', 'Showed one repeated value in every quarter in the 8/31 file (date pointed at a blank cell) - fixed'),
        ('Balance Sheet Evolution', 'TBVPS', 'Showed one repeated value in every quarter in the 8/31 file - fixed'),
        ]:
    add(FIXED, sheet, metric, why=why, value='')
add(FIXED, 'Valuation', 'P/E', 'KEY', ['2024/3F'], why='Client cell W16 (1,523x) now displays NM', value='NM')

# ------------------------------------------------------------------ sheet
if 'Data Check' in wb.sheetnames:
    del wb['Data Check']
ws = wb.create_sheet('Data Check')
ws.sheet_view.showGridLines = False
navy = PatternFill('solid', fgColor='1F4E78')
ws['A1'] = 'Data Check - values that are probably wrong or worth a look'
ws['A1'].font = Font(name='Arial', bold=True, size=13, color='FFFFFF')
ws['A1'].fill = navy
ws.merge_cells('A1:I1')
ws.row_dimensions[1].height = 22
ws['A2'] = ('Checked against the values FactSet returned in your 8/31 refresh, shown the way this file displays them. '
            'Click a cell reference to jump to it. Internal - remove this sheet and Scratch Notes before sending to the client.')
ws['A2'].font = Font(name='Arial', italic=True, size=9, color='FF595959')
ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells('A2:I2')
ws.row_dimensions[2].height = 28
hdr = ['Status', 'Worksheet', 'Metric', 'Bank', 'Quarter(s)', 'Cell', 'Value shown', "What's wrong", 'Likely cause / next step']
thin = Side(style='thin', color='D9D9D9')
brd = Border(left=thin, right=thin, top=thin, bottom=thin)
for i, h in enumerate(hdr, 1):
    c = ws.cell(row=4, column=i, value=h)
    c.font = Font(name='Arial', bold=True, size=10, color='FFFFFF')
    c.fill = navy
    c.border = brd
    c.alignment = Alignment(horizontal='center', vertical='center')
FILL = {PROB: 'F8D7DA', CHECK: 'FFF3CD', REAL: 'E2EFDA', FIXED: 'DDEBF7'}
r = 5
for row in rows:
    for j, v in enumerate(row, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.font = Font(name='Arial', size=9, bold=(j == 1))
        c.border = brd
        c.alignment = Alignment(vertical='top', wrap_text=j in (7, 8, 9))
    ws.cell(row=r, column=1).fill = PatternFill('solid', fgColor=FILL[row[0]])
    link = ws.cell(row=r, column=6)
    link.hyperlink = Hyperlink(ref=link.coordinate, location=f"'{row[1]}'!{row[5]}", display=row[5])
    link.font = Font(name='Arial', size=9, color='0563C1', underline='single')
    r += 1
for col, w in zip('ABCDEFGHI', (17, 22, 26, 7, 17, 7, 20, 48, 58)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A5'
wb.save(OUT)

from collections import Counter
print('saved', OUT)
print(Counter(r[0] for r in rows))

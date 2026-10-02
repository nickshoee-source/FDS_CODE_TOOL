"""
V9 formatting pass on top of V8.

 1. Number formats per metric (metric_formats.py): $ with commas, whole-number
    % with a % sign, multiples with x, P/E shows NM when >100x or negative,
    negatives in parentheses, thousands separators everywhere.
 2. Percent convention unified to whole numbers (FactSet's native form):
      - TSR: dropped the /100 inside the FactSet formula
      - Buyback Yield, Capital Return Yield: cell math x100
      - historical Efficiency Ratio (FF_EFF_RATIO returns 0.605): x100
 3. #N/A -> "-": every historical per-cell FactSet pull wrapped in
    _xlfn.IFNA(...,"-"). (IFNA is an Excel-2013 function, so the file format
    requires the _xlfn. prefix -- without it Excel shows #NAME?; verified.)
    IFNA only catches #N/A, so genuine errors (#VALUE!, #DIV/0!) stay visible.
 4. Estimates sheets, metrics with no FE time series: a date-conditional IF --
    "-" while the quarter is still in the future, otherwise the reported value
    (so the sheet fills in by itself as 2027+ quarters get reported, and no
    FactSet call is made for future quarters).
 5. Conditional formatting (client request): "-" and #N/A cells in grey.
 6. Units note in row 2 of each sheet that carries $ values.

Skipped on purpose: company-name pulls (all resolve), FDSRC time-series anchors
(already default to "-"), the user's Valuation 2026/3F-4F spill arrays.
"""
import re
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.styles import Font
from openpyxl.formatting.rule import FormulaRule
from metric_formats import FMT, KIND, PER_SHARE, base_sheet

SRC = '/home/claude/Quarterly_Template_V8.xlsx'
OUT = '/home/claude/Quarterly_Template_V9.xlsx'
wb = load_workbook(SRC, data_only=False)


def txt(c):
    v = c.value
    return v.text if isinstance(v, ArrayFormula) else v


def s_(v):
    return v if isinstance(v, str) else ''


def C(i):
    return get_column_letter(i)


def runs(ws, mr):
    out, cur, st = [], None, None
    for c in range(5, ws.max_column + 2):
        m = txt(ws.cell(row=mr, column=c))
        if not isinstance(m, str):
            m = None
        if m != cur:
            if cur is not None:
                out.append((cur, st, c - 1))
            cur, st = m, c
    return out


PREF_RE = re.compile(r'(?<![A-Z])\$?[A-Z]{1,3}\$\d+')

stats = {k: 0 for k in ['fmt', 'ifna', 'est_if', 'x100', 'tsr', 'cf_ranges', 'notes']}
GREY = Font(color='FF9E9E9E')

for sheet in wb.sheetnames:
    if sheet == 'Scratch Notes':
        continue
    ws = wb[sheet]
    base = base_sheet(sheet)
    is_est = sheet.endswith('- Estimates')
    NP = 12 if is_est else 28

    # multi-cell array spills (user's Valuation 2026/3F-4F) -- leave formulas alone
    spill_cells = set()
    for row in ws.iter_rows():
        for cl in row:
            if isinstance(cl.value, ArrayFormula) and ':' in cl.value.ref:
                a, b = cl.value.ref.split(':')
                if a != b:
                    ma, mb = re.match(r'([A-Z]+)(\d+)', a), re.match(r'([A-Z]+)(\d+)', b)
                    for rr in range(int(ma.group(2)), int(mb.group(2)) + 1):
                        for cc in range(column_index_from_string(ma.group(1)),
                                        column_index_from_string(mb.group(1)) + 1):
                            spill_cells.add((rr, cc))

    has_dollar, has_shares, has_ps = False, False, False
    for cr in (3, 22):
        if txt(ws.cell(row=cr + 1, column=5)) is None:
            continue
        first, last = cr + 5, cr + 15
        block_last_col = 4
        for lab, s, e in runs(ws, cr + 1):
            block_last_col = max(block_last_col, e)
            kind = KIND[(base, lab)]
            fmt = FMT[kind]
            if kind == 'DOLLAR':
                if (base, lab) in PER_SHARE:
                    has_ps = True
                else:
                    has_dollar = True
            if kind == 'NUM':
                has_shares = True
            f_first = s_(txt(ws.cell(row=first, column=s)))
            is_ts_block = 'FE_TIMESERIES' in f_first and f_first.startswith('=FDSRC(')

            for i in range(NP):
                col = s + i
                for r in range(first, last + 1):
                    cell = ws.cell(row=r, column=col)
                    cell.number_format = fmt
                    stats['fmt'] += 1
                    if is_ts_block or (r, col) in spill_cells:
                        continue
                    f = s_(txt(cell))
                    if not f.startswith('='):
                        continue

                    # ---- cell math: whole-number percent for the two yields
                    if f.upper().startswith('=IFERROR('):
                        if lab in ('Buyback Yield', 'Capital Return Yield') and '*100,"-")' not in f:
                            inner = f[len('=IFERROR('):-len(',"-")')]
                            cell.value = f'=IFERROR(({inner})*100,"-")'
                            stats['x100'] += 1
                        continue

                    if 'FDS' not in f.upper():
                        continue

                    inner = f[1:].replace('_xll.', '')
                    # ---- TSR: whole-number percent (drop the /100)
                    if lab == 'TSR' and ')/100")' in inner:
                        inner = inner.replace(')/100")', ')")')
                        stats['tsr'] += 1
                    # ---- historical Efficiency Ratio returns a fraction
                    if lab == 'Efficiency Ratio' and not is_est:
                        inner = f'{inner}*100'
                        stats['x100'] += 1

                    if is_est:
                        p = PREF_RE.search(inner).group(0)          # the cell's own period cell
                        cell.value = (f'=IF(DATE(VALUE(LEFT({p},4)),VALUE(MID({p},6,1))*3+1,0)'
                                      f'>TODAY(),"-",_xlfn.IFNA({inner},"-"))')
                        stats['est_if'] += 1
                    else:
                        cell.value = f'=_xlfn.IFNA({inner},"-")'
                        stats['ifna'] += 1

        # ---- conditional formatting: "-" and #N/A in grey
        rng = f'E{first}:{C(block_last_col)}{last}'
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'IFERROR(E{first}="-",ISNA(E{first}))'], font=GREY))
        stats['cf_ranges'] += 1

    # ---- units note (row 2)
    parts = []
    if has_dollar:
        parts.append('$ in millions' + (', except per-share data' if has_ps else ''))
    elif has_ps:
        parts.append('$ per share')
    if has_shares:
        parts.append('shares in millions')
    if parts:
        note = ws.cell(row=2, column=1)
        note.value = '; '.join(parts).capitalize().replace('$ in', '$ in') + '.'
        note.value = note.value[0].upper() + note.value[1:]
        note.font = Font(name='Arial', italic=True, size=9, color='FF595959')
        stats['notes'] += 1

wb.save(OUT)
print('saved', OUT)
for k, v in stats.items():
    print(f'  {k:10} {v}')

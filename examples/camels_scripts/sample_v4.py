from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter as L
SRC='/mnt/user-data/outputs/CAMELS_FDS_Code_Analysis_V3.xlsx'
OUT='/mnt/user-data/outputs/CAMELS_FDS_Code_Analysis_V4_SAMPLE_5_Banks.xlsx'
wb=load_workbook(SRC); H=wb['CAMELS HISTORICAL ANNUAL']
YEL=PatternFill('solid',fgColor='FFFFFF00')
I=lambda c,u=False:f'{c}(ANN_L,{{Y}}'+(',,,,USD)' if u else ')')   # FFI
F=lambda c,u=False:f'{c}(ANN_R,{{Y}}'+(',,,,USD)' if u else ')')   # FF
B=lambda c,u=False:f'{c}(ANN,{{Y}},,,RF'+(',USD)' if u else ')')   # FB (regulatory, US banks)
ch=lambda *x:'@'.join(x)
LOANS=ch(I('FFI_LOAN_ADV_TOT',1),F('FF_BK_LOAN_TOT',1),B('FB_TOT_HFI_HFS_UNEARN_INC',1))
DEPS=ch(I('FFI_DEPS_TOT',1),F('FF_DEPS',1),B('FB_DEPS',1))
CODES={
 6: ch(I('FFI_COM_EQ_TIER1_TOT',1),F('FF_BK_COM_EQ_TIER1_TOT',1),B('FB_COM_EQ_TIER1',1)),
 7: ch(I('FFI_COM_EQ_TIER1_RATIO'),F('FF_BK_COM_EQ_TIER1_RATIO'),B('FB_COM_EQ_TIER1_RATIO')),
 8: ch(I('FFI_TIER1_CAP',1),F('FF_TIER1_CAP',1),B('FB_TIER1_CAP',1)),
 9: ch(I('FFI_CAP_RATIO_TIER1'),F('FF_CAP_RATIO_TIER1'),B('FB_CAP_RATIO_TIER1')),
 10: ch(I('FFI_CAP_RATIO_TOT'),F('FF_CAP_RATIO_TOT'),B('FB_CAP_RATIO_TOT')),
 11: ch(I('FFI_LEV_RATIO_RPT'),F('FF_BK_LEV_RATIO'),B('FB_LEV_RATIO')),
 12: ch(I('FFI_RWA',1),F('FF_ASSETS_RISK_WGHT',1),B('FB_ASSETS_RISK_WGHT',1)),
 14: ch(I('FFI_NPL_LOAN_ADV',1),F('FF_LOAN_NONPERF',1),B('FB_NONPERF_LOAN',1)),
 15: ch(I('FFI_LOAN_LOSS_PROV',1),F('FF_LOAN_LOSS_PROV',1),B('FB_LOAN_LOSS_PROV',1)),
 16: '('+ch(I('FFI_NPL_LOAN_LOSS_RSRV_RATIO'),F('FF_NONPERF_LOAN_LOSS_RSRV'))+')/100',
 17: '('+ch(I('FFI_NPL_LOAN_RATIO'),F('FF_NONPERF_LOAN_PCT'),B('FB_NONPERF_LOAN_PCT'))+')/100',
 19: ch(I('FFI_EFF_RATIO',1),F('FF_BK_EFF_RATIO'),B('FB_EFF_RATIO')),
 21: ch(I('FFI_ROTE'),B('FB_ROTE')),
 22: ch(I('FFI_ROTCE'),F('FF_ROTCE'),B('FB_ROTCE')),
 23: ch(I('FFI_ROA'),F('FF_ROA'),B('FB_ROA')),
 24: ch(I('FFI_AVG_BAL_INT_RATE_NET_MGN'),F('FF_INT_MGN'),B('FB_INT_MGN')),
 25: ch(I('FFI_INT_INC_NET',1),F('FF_INT_INC_NET',1),B('FB_INT_INC_NET',1)),
 26: ch(I('FFI_NON_INT_INC',1),F('FF_NON_INT_INC',1),B('FB_NON_INT_INC',1)),
 28: LOANS, 29: DEPS, 30: f'({LOANS})/({DEPS})',
 31: ch(I('FFI_BK_LIQ_COVG_RATIO'),F('FF_BK_LIQ_COVG_RATIO')),
 32: I('FFI_BK_NSFR'),
 34: ch(I('FFI_LOAN_AMORT_CUST',1),F('FF_LOAN_NET',1),B('FB_TOT_HFI_HFS_UNEARN_INC',1)), 35: I('FFI_LOAN_AMORT_BK',1),
 36: '('+I('FFI_SECS_INVEST',1)+'+'+I('FFI_TRADE_ACCT',1)+'+'+I('FFI_DERIV_HEDGE',1)+')@'+F('FF_INVEST_TOT',1),
 37: '('+ch(I('FFI_DEPS',1),F('FF_DEPS_CUST',1))+')/('+ch(I('FFI_DEPS_TOT',1),F('FF_DEPS',1),B('FB_DEPS',1))+')',
}
for c in range(3,8):
    col=L(c)
    for r,code in CODES.items():
        H.cell(r,c).value=f'=FDSC("-",{col}2,"'+code.replace('{Y}',f'"&{col}4&"')+'")'
        H.cell(r,c).fill=YEL
    H.cell(38,c).value=f'=IFERROR(1-{col}37,"-")'
wb.save(OUT)
for r,code in CODES.items(): print(r,H.cell(r,2).value,'|',H.cell(r,3).value)

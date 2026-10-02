"""
fds_kit.py - reusable helpers for building / editing FactSet =FDS workbooks
with openpyxl.  Import from scripts:   from tools.fds_kit import *

Sections
  1. Formula builders   (FDS / FDSC / FDSRC strings, fallback chains, wrappers)
  2. Styles             (Quarterly-Template blue, red, number formats)
  3. Change tracking    (yellow-highlight exactly the changed cells, change log)
  4. Workbook hygiene   (freeze panes, sensitivity-label restore)
  5. Verification       (LibreOffice parse check, code harvest from refreshed files)

Rules baked in (see docs/FactSet_FDS_Coding_Guide.md):
  * write plain '=FDS(...)' strings - never ArrayFormula / '_xll.' (causes {} and @)
  * Excel 2013+ functions need the _xlfn. prefix when written by openpyxl (_xlfn.IFNA)
  * yellow FFFF00 means "changed from the original" and nothing else
"""
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

# =============================================================== 1. formulas

def ref(cell: str) -> str:
    """Splice a cell into an FQL string:  ref('E$5') -> '"&E$5&"'"""
    return f'"&{cell}&"'


def fds(id_cell: str, fql: str) -> str:
    return f'=FDS({id_cell},"{fql}")'


def fdsc(id_cell: str, fql: str, default: str = '-') -> str:
    return f'=FDSC("{default}",{id_cell},"{fql}")'


def fdsrc(id_cell: str, fql: str, default: str = '-') -> str:
    """Row spill (FE_TIMESERIES). Write only in the first cell; leave cells to the right empty."""
    return f'=FDSRC("{default}",{id_cell},"{fql}")'


def ifna(formula: str, default: str = '-') -> str:
    """=_xlfn.IFNA(FDS(...),"-")  - strips the leading '=' of the inner formula."""
    return f'=_xlfn.IFNA({formula.lstrip("=")},"{default}")'


def future_gate(period_cell: str, formula: str) -> str:
    """'-' while a 'YYYY/QF' quarter is still in the future, else the IFNA-wrapped pull."""
    p = period_cell
    inner = formula.lstrip('=')
    if not inner.startswith('_xlfn.IFNA('):
        inner = f'_xlfn.IFNA({inner},"-")'
    return f'=IF(DATE(VALUE(LEFT({p},4)),VALUE(MID({p},6,1))*3+1,0)>TODAY(),"-",{inner})'


def chain(*items: str) -> str:
    """FQL fallback: chain(a,b,c) -> 'a@b@c'  (use b if a is NA, then c). Test before rollout."""
    return '@'.join(i for i in items if i)


# periodicity helpers -> return the FQL item text with a period placeholder already spliced
def ff(item, period, per='ANN_R', usd=False):
    """FactSet Fundamentals: FF_X(ANN_R,2025[,,,,USD])"""
    return f'{item}({per},{period}' + (',,,,USD)' if usd else ')')


def ffi(item, period, per='ANN_L', usd=False):
    """FactSet Fundamentals Industry: FFI_X(ANN_L,2025[,,,,USD]) - USD or you get local currency"""
    return f'{item}({per},{period}' + (',,,,USD)' if usd else ')')


def fb(item, period, per='ANN', usd=False):
    """FactSet regulatory (US banks): FB_X(ANN,2025,,,RF[,USD]);  quarterly: fb(item,p,'QTR') -> FB_X(QTR,p)"""
    if per == 'QTR':
        return f'{item}(QTR,{period})'
    return f'{item}({per},{period},,,RF' + (',USD)' if usd else ')')


def fe_timeseries(item, start_cell, end_cell, valuation=False, freq='FQ'):
    """FE_TIMESERIES / FE_TIMESERIES_VALUATION with the HBAN option strings."""
    if valuation:
        return (f"FE_TIMESERIES_VALUATION({item},MEAN,{ref(start_cell)},{ref(end_cell)},{freq},"
                f"'BKRACTMED=1,WIN=0,UNITS=AUTO,DATE=NOW,CALC=LTMA')")
    return (f"FE_TIMESERIES({item},MEAN,{ref(start_cell)},{ref(end_cell)},{freq},"
            f"'BKRACTMED=1,WIN=0,CURRENCY=RPT,UNITS=AUTO,DATE=NOW')")


def fe_estimate(item, per='NTMA', period='', currency='USD'):
    """FE_ESTIMATE(ITEM,MEAN,NTMA,,NOW,,,'CURRENCY=USD') ; per='ANN_ROLL', period=ref('C4') for a fiscal year."""
    opt = f"'CURRENCY={currency}'" if currency else "''"
    return f'FE_ESTIMATE({item},MEAN,{per},{period},NOW,,,{opt})'


# =============================================================== 2. styles

YELLOW = PatternFill('solid', fgColor='FFFFFF00')
NOFILL = PatternFill(fill_type=None)

FMT = {
    'PCT':      '#,##0.00"%"_);(#,##0.00"%")',      # value already a whole-number %
    'PCT_FRAC': '0.00%_);(0.00%)',                  # value is a fraction
    'DOLLAR':   '$#,##0.00_);($#,##0.00)',
    'DPS':      '$#,##0.00#_);($#,##0.00#)',
    'MULT':     '[<0]"NM";#,##0.00"x"',
    'PE':       '[>100]"NM";[<0]"NM";#,##0.00"x"',
    'NUM':      '#,##0.00_);(#,##0.00)',
    'COUNT':    '#,##0_);(#,##0)',
}

THEMES = {
    'quarterly': dict(header_fill='FFD9E7F5', header_font='FF1F4E78', section_fill='FFD9E7F5',
                      section_font='FF1F4E78', period_fill='FFF2F2F2'),
    'red':      dict(header_fill='FFB00B1C', header_font='FFFFFFFF', section_fill='FFF7E1E4',
                      section_font='FFB00B1C', period_fill='FFF7E1E4'),
}


def style_header(cells, theme='quarterly', size=9):
    t = THEMES[theme]
    for c in cells:
        c.fill = PatternFill('solid', fgColor=t['header_fill'])
        c.font = Font(name='Arial', size=size, bold=True, color=t['header_font'])
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def style_section(cells, theme='quarterly'):
    t = THEMES[theme]
    for c in cells:
        c.fill = PatternFill('solid', fgColor=t['section_fill'])
        c.font = Font(name='Arial', size=10, bold=True, color=t['section_font'])


def style_period_row(cells, theme='quarterly'):
    t = THEMES[theme]
    line = Border(bottom=Side(style='medium', color=t['section_font'])) if theme == 'red' else Border()
    for c in cells:
        c.fill = PatternFill('solid', fgColor=t['period_fill'])
        c.font = Font(name='Arial', size=8 if theme == 'quarterly' else 9, bold=theme == 'red',
                      color='FF595959' if theme == 'quarterly' else t['section_font'])
        c.alignment = Alignment(horizontal='center')
        c.border = line


def grey_dash_rule(ws, rng: str):
    """Grey font for '-' and #N/A cells. rng like 'E8:AF18' (rule anchored on the top-left cell)."""
    tl = rng.split(':')[0]
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'IFERROR({tl}="-",ISNA({tl}))'],
                                                   font=Font(color='FF9E9E9E')))


# =============================================================== 3. change tracking

def formula_text(v):
    """Normalise a cell value for comparison (ArrayFormula -> text, drop _xll.)."""
    v = v.text if isinstance(v, ArrayFormula) else v
    return v.replace('_xll.', '') if isinstance(v, str) else v


def is_yellow(cell) -> bool:
    f = cell.fill
    return f.fill_type == 'solid' and f.fgColor.type == 'rgb' and f.fgColor.rgb == 'FFFFFF00'


def highlight_changes(original_path, wb, sheets=None, remove_other_yellow=True):
    """Make yellow == 'content differs from the original file' on every shared sheet.
    Returns {sheet: (changed, yellow_added, yellow_removed)}."""
    orig = load_workbook(original_path)
    stats = {}
    for name in (sheets or orig.sheetnames):
        if name not in wb.sheetnames:
            continue
        ws, ow = wb[name], orig[name]
        changed = added = removed = 0
        for r in range(1, max(ws.max_row, ow.max_row) + 1):
            for c in range(1, max(ws.max_column, ow.max_column) + 1):
                cell = ws.cell(r, c)
                diff = formula_text(cell.value) != formula_text(ow.cell(r, c).value)
                changed += diff
                if diff and not is_yellow(cell):
                    cell.fill = YELLOW
                    added += 1
                elif not diff and is_yellow(cell) and remove_other_yellow:
                    cell.fill = NOFILL
                    removed += 1
        stats[name] = (changed, added, removed)
    return stats


def write_change_log(wb, rows, title='Change Log', theme='red'):
    """rows: iterable of [sheet, row, metric, cells, before, after, why, status]."""
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    ws.append(['Sheet', 'Row', 'Metric', 'Cells', 'Before (first cell)', 'After (first cell)', 'Why', 'Status'])
    for r in rows:
        ws.append(list(r))
    style_header(ws[1], theme)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name='Arial', size=9)
            c.alignment = Alignment(wrap_text=True, vertical='top')
        if 'CONFIRM' in str(row[7].value).upper():
            row[7].font = Font(name='Arial', size=9, bold=True, color='FFB00B1C')
    for col, w in zip('ABCDEFGH', (26, 6, 34, 14, 60, 60, 60, 28)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    return ws


# =============================================================== 4. hygiene

def freeze(ws, cell: str):
    """Keep labels/IDs visible while scrolling (HBAN 'E8', CAMELS formatted 'D3', CAMELS source 'C5')."""
    ws.freeze_panes = cell
    ws.sheet_view.topLeftCell = 'A1'


def even_widths(ws, first_col: int, last_col: int, width: float = 14):
    for c in range(first_col, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = width


def restore_sensitivity_label(original_path, saved_path):
    """openpyxl drops docMetadata/LabelInfo.xml (Microsoft sensitivity label). Copy it back."""
    with zipfile.ZipFile(original_path) as z:
        if 'docMetadata/LabelInfo.xml' not in z.namelist():
            return False
        label = z.read('docMetadata/LabelInfo.xml')
    tmp = str(saved_path) + '.tmp'
    with zipfile.ZipFile(saved_path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == '[Content_Types].xml' and b'LabelInfo' not in data:
                data = data.replace(b'</Types>', b'<Override PartName="/docMetadata/LabelInfo.xml" '
                                    b'ContentType="application/vnd.ms-office.classificationlabels+xml"/></Types>')
            if item.filename == '_rels/.rels' and b'LabelInfo' not in data:
                data = data.replace(b'</Relationships>', b'<Relationship Type="http://schemas.microsoft.com/'
                                    b'office/2020/02/relationships/classificationlabels" '
                                    b'Target="docMetadata/LabelInfo.xml" Id="rIdLbl1"/></Relationships>')
            if item.filename != 'docMetadata/LabelInfo.xml':
                zout.writestr(item, data)
        zout.writestr('docMetadata/LabelInfo.xml', label)
    shutil.move(tmp, saved_path)
    return True


# =============================================================== 5. verification

def libreoffice_parse_check(path):
    """Convert with LibreOffice (if installed) and return formula parse errors (Err:5xx).
    FactSet functions show #NAME? outside Excel - that is expected and not reported."""
    exe = shutil.which('soffice') or shutil.which('libreoffice')
    if not exe:
        return None, 'LibreOffice not installed - skipped (install it for this check, optional)'
    with tempfile.TemporaryDirectory() as d:
        subprocess.run([exe, '--headless', '--convert-to', 'ods', '--outdir', d, str(path)],
                       capture_output=True, timeout=600)
        ods = Path(d) / (Path(path).stem + '.ods')
        if not ods.exists():
            return None, 'conversion failed'
        xml = zipfile.ZipFile(ods).read('content.xml').decode('utf8', 'ignore')
        errs = sorted(set(re.findall(r'Err:\d+', xml)))
        return errs, 'ok' if not errs else 'PARSE ERRORS FOUND'


CODE_RX = re.compile(r'\b((?:FFI|FF|FB|FG|FMA|FREF|P)_[A-Z0-9_]+)\s*\(')
FE_ITEM_RX = re.compile(r'FE_[A-Z_]+\(\s*([A-Z0-9_]+)\s*,')


def harvest_codes(refreshed_path):
    """Read a workbook that was REFRESHED in Excel+FactSet and report, per code, how many
    cells returned numbers / NA / text. This is the evidence that a code works."""
    F = load_workbook(refreshed_path)
    V = load_workbook(refreshed_path, data_only=True)
    out = defaultdict(lambda: {'cells': 0, 'numbers': 0, 'na_or_error': 0, 'text': 0,
                               'samples': [], 'example': None})
    for sh in F.sheetnames:
        for row in F[sh].iter_rows():
            for c in row:
                t = formula_text(c.value)
                if not (isinstance(t, str) and t.startswith('=') and 'FDS' in t.upper()):
                    continue
                cv = V[sh][c.coordinate].value
                keys = set(CODE_RX.findall(t)) | {'FE:' + m for m in FE_ITEM_RX.findall(t)}
                for k in keys:
                    d = out[k]
                    d['cells'] += 1
                    d['example'] = d['example'] or f'{sh}!{c.coordinate}: {t}'
                    if isinstance(cv, (int, float)):
                        d['numbers'] += 1
                        if len(d['samples']) < 3:
                            d['samples'].append(cv)
                    elif cv is None or (isinstance(cv, str) and cv.startswith(('#', 'ERROR'))):
                        d['na_or_error'] += 1
                    else:
                        d['text'] += 1
    return dict(out)

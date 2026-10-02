"""
V7: convert the "- Estimates" worksheets to forward-looking FactSet
FE_TIMESERIES / FE_TIMESERIES_VALUATION pulls, using the item codes from the
user's All_Estimates_Formulas.xlsx (HBAN-US) template.

Mechanics, matching the user's template exactly:
  * The template's working formulas are FDSRC("-", <ticker>, "FE_TIMESERIES(...)")
    entered in ONE cell and spilling right across the period range (the raw XML
    shows <f t="array" ref="B6:L6"> on the anchor, with the trailing cells
    holding only cached values). FDS cannot spill -- FDSRC is required.
  * So each metric gets ONE formula in its FIRST period column per bank row,
    and the remaining 11 period cells are left EMPTY for FactSet to fill on
    refresh. That is exactly the "put it in the first cell and it autopopulates"
    behaviour the user described, and it avoids re-introducing the {} / @
    artifacts that a hand-written CSE array caused earlier in this project.
  * The date range is driven by the sheet's own period row, so the formula stays
    refreshable: "&<first>$<prow>&" .. "&<last>$<prow>&" -> 2027/1F .. 2029/4F.

Also fixed here (pre-existing bug, both the historical and Estimates sheet):
  Profitability's 'Effective Tax Rate' is the 4th metric of the LOWER block, but
  its formula read the period from row 5 -- the UPPER block only has 3 metrics,
  so that cell is blank. Lower-block formulas now read their own period row.
"""
import re
import copy
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter, column_index_from_string

SRC = '/home/claude/Quarterly_Template_V6.xlsx'
OUT = '/home/claude/Quarterly_Template_V7.xlsx'
N_EST = 12

wb = load_workbook(SRC, data_only=False)


def txt(c):
    v = c.value
    return v.text if isinstance(v, ArrayFormula) else v


def s_(v):
    return v if isinstance(v, str) else ''


def C(i):
    return get_column_letter(i)


def metric_runs(ws, met_row):
    out, cur, st = [], None, None
    for c in range(5, ws.max_column + 2):
        m = txt(ws.cell(row=met_row, column=c))
        if not isinstance(m, str):
            m = None
        if m != cur:
            if cur is not None:
                out.append((cur, st, c - 1))
            cur, st = m, c
    return out


# ---- option strings lifted verbatim from the user's template -------------
OPT_TS = "'BKRACTMED=1,WIN=0,CURRENCY=RPT,UNITS=AUTO,DATE=NOW'"
OPT_VAL = "'BKRACTMED=1,WIN=0,UNITS=AUTO,DATE=NOW,CALC=LTMA'"

TS, VAL = 'TS', 'VAL'

# ---- metric -> (function, FactSet estimate item) ------------------------
MAP = {
    ('Valuation - Estimates', 'P/E'):                          (VAL, 'PE'),
    ('Valuation - Estimates', 'P/B'):                          (VAL, 'PBPS'),
    ('Valuation - Estimates', 'P/TBV'):                        (VAL, 'P_TBPS'),

    ('Shareholder Returns - Estimates', 'Dividend Yield'):      (VAL, 'DIV_YLD'),
    ('Shareholder Returns - Estimates', 'Share Repurchases'):   (TS, 'SHS_REPURCH'),
    ('Shareholder Returns - Estimates', 'Dividend Per Share'):  (TS, 'NETDIV'),
    ('Shareholder Returns - Estimates', 'Stock Repurchase CF'): (TS, 'SHS_REPURCH'),

    ('Profitability - Estimates', 'ROTCE'):                     (TS, 'ROTE'),
    ('Profitability - Estimates', 'ROE'):                       (VAL, 'ROE'),
    ('Profitability - Estimates', 'ROA'):                       (VAL, 'ROA'),
    ('Profitability - Estimates', 'EPS'):                       (TS, 'EPS'),
    ('Profitability - Estimates', 'Net Income'):                (TS, 'NETPROFIT'),
    ('Profitability - Estimates', 'PPNR'):                      (TS, 'EBIT_ADJ'),

    ('Revenue & Fee Mix - Estimates', 'NII'):                   (TS, 'NETINTERESTINC'),
    ('Revenue & Fee Mix - Estimates', 'Non-Interest Income'):   (TS, 'NON_INT_INC'),

    ('Efficiency - Estimates', 'Efficiency Ratio'):             (TS, 'COST_INCOME'),
    ('Efficiency - Estimates', 'NIE'):                          (TS, 'OPEREXPEN'),
    ('Efficiency - Estimates', 'Personnel Expense'):            (TS, 'SAL_BENEFITS'),

    ('NIM & Margin - Estimates', 'NIM'):                        (TS, 'INT_INC_MARGIN'),

    ('Balance Sheet Evol - Estimates', 'Total Assets EOP'):     (TS, 'TOTASSET'),
    ('Balance Sheet Evol - Estimates', 'Intangible Assets'):    (TS, 'INTANG'),
    ('Balance Sheet Evol - Estimates', 'Total Loans EOP'):      (TS, 'LOAN_GROSS'),
    ('Balance Sheet Evol - Estimates', 'Earning Assets EOP'):   (TS, 'AVG_EARN_ASSETS'),
    ('Balance Sheet Evol - Estimates', 'Total Deposits EOP'):   (TS, 'DEPS'),
    ('Balance Sheet Evol - Estimates', 'TBVPS'):                (TS, 'BPS_TANG'),

    ('Deposit Competition - Estimates', 'NIBD Balance'):            (TS, 'AVG_NINTB_DEPS'),
    ('Deposit Competition - Estimates', 'Interest Bearing Deposits'): (TS, 'AVG_INTB_DEPS'),
    ('Deposit Competition - Estimates', 'Total Deposits'):          (TS, 'DEPS'),

    ('Credit - Estimates', 'NCO Ratio'):                        (VAL, 'NETCHARGE_LOANNET'),
    ('Credit - Estimates', 'NPL Ratio'):                        (VAL, 'NPL'),
    ('Credit - Estimates', 'NPA Ratio'):                        (TS, 'LOANOREO_NONPERF'),
    ('Credit - Estimates', 'ACL Ratio'):                        (TS, 'LOANLOSSRSV_PCT'),

    ('Capital - Estimates', 'CET1 Ratio'):                      (TS, 'COMCAP_RATIO_TIER1'),
    ('Capital - Estimates', 'Risk Weighted Assets'):            (TS, 'ASSETS_RISK_WGHT'),

    ('Estimates - Estimates', 'Consensus EPS'):                 (TS, 'CUSTOM_EPS'),
    ('Estimates - Estimates', 'Consensus Revenue'):             (TS, 'SALES'),
    ('Estimates - Estimates', 'Consensus NII'):                 (TS, 'NETINTERESTINC'),
    ('Estimates - Estimates', 'Consensus Fee Income'):          (TS, 'INCFEESCOM'),
    ('Estimates - Estimates', 'Consensus ROTCE'):               (TS, 'ROTE'),
    ('Estimates - Estimates', 'Consensus CET1'):                (TS, 'COM_EQUITY_TIER1'),
}

# metric mapped to an item whose basis differs from the historical sheet
BASIS_CHANGED = {
    ('Balance Sheet Evol - Estimates', 'Earning Assets EOP'):
        'period-end -> consensus AVERAGE earning assets (no EOP estimate exists)',
    ('Deposit Competition - Estimates', 'NIBD Balance'):
        'period-end -> consensus AVERAGE non-interest-bearing deposits',
    ('Deposit Competition - Estimates', 'Interest Bearing Deposits'):
        'period-end -> consensus AVERAGE interest-bearing deposits',
    ('Balance Sheet Evol - Estimates', 'Total Loans EOP'):
        'uses GROSS loans (LOAN_GROSS); LOAN_NET is also available',
    ('Credit - Estimates', 'NPL Ratio'):
        "item NPL is labelled 'NPLs and Net Charge-Offs/Average Loans' in your "
        "template - worth confirming it is the NPL ratio you want",
}

# left as native cell-reference math (do not overwrite)
CELLMATH = {
    ('Shareholder Returns - Estimates', 'Buyback Yield'),
    ('Shareholder Returns - Estimates', 'Capital Return Yield'),
    ('Revenue & Fee Mix - Estimates', 'Total Revenue'),
    ('NIM & Margin - Estimates', 'Cost of Funds'),
    ('Balance Sheet Evol - Estimates', 'TBV'),
}


def build(func, item, row, first_col, last_col, prow):
    fn = 'FE_TIMESERIES' if func == TS else 'FE_TIMESERIES_VALUATION'
    opt = OPT_TS if func == TS else OPT_VAL
    return (f'=FDSRC("-",$D{row},"{fn}({item},MEAN,"&{C(first_col)}${prow}'
            f'&","&{C(last_col)}${prow}&",FQ,{opt})")')


replaced, unmapped, kept = [], [], []

for sheet in [n for n in wb.sheetnames if n.endswith('- Estimates')]:
    ws = wb[sheet]
    for cat_row in (3, 22):
        met_row, per_row, first = cat_row + 1, cat_row + 2, cat_row + 5
        if txt(ws.cell(row=met_row, column=5)) is None:
            continue
        last_row = first + 10
        for lab, s, e in metric_runs(ws, met_row):
            key = (sheet, lab)
            if key in CELLMATH:
                kept.append(key)
                continue
            if key not in MAP:
                unmapped.append((sheet, lab, s_(txt(ws.cell(row=first, column=s)))))
                continue
            func, item = MAP[key]
            for r in range(first, last_row + 1):
                # anchor formula in the first period column
                a = ws.cell(row=r, column=s)
                a.value = build(func, item, r, s, s + N_EST - 1, per_row)
                a.number_format = '0.00'
                # clear the remaining period cells so FactSet can spill into them
                for i in range(1, N_EST):
                    c = ws.cell(row=r, column=s + i)
                    c.value = None
                    c.number_format = '0.00'
            replaced.append((sheet, lab, func, item, C(s)))

# ---------- fix the lower-block period-row bug on Profitability ----------
fixes = []
for sheet in ['Profitability', 'Profitability - Estimates']:
    ws = wb[sheet]
    cat_row, met_row, per_row, first = 22, 23, 24, 27
    upper_last = max(e for _, _, e in metric_runs(ws, 4))
    for lab, s, e in metric_runs(ws, met_row):
        if s <= upper_last:
            continue                      # its row-5 reference is populated
        for r in range(first, first + 11):
            cell = ws.cell(row=r, column=s)
            f = s_(txt(cell))
            if not f:
                continue
            new = re.sub(r'([A-Z]{1,3})\$5\b', r'\1$%d' % per_row, f).replace('_xll.', '')
            if new != f:
                cell.value = new
        # and across its remaining period columns (historical sheet only)
        for i in range(1, 28):
            cc = ws.cell(row=first, column=s + i)
            if s_(txt(cc)) == '':
                break
            for r in range(first, first + 11):
                cell = ws.cell(row=r, column=s + i)
                f = s_(txt(cell))
                if f:
                    cell.value = re.sub(r'([A-Z]{1,3})\$5\b',
                                        r'\1$%d' % per_row, f).replace('_xll.', '')
        fixes.append((sheet, lab, C(s)))

wb.save(OUT)

print('saved', OUT)
print(f'\n{"="*92}\nREPLACED with forward-looking time-series pulls ({len(replaced)})\n{"="*92}')
cur = None
for sheet, lab, func, item, col in replaced:
    if sheet != cur:
        print(f'\n  {sheet}')
        cur = sheet
    tag = 'FE_TIMESERIES          ' if func == TS else 'FE_TIMESERIES_VALUATION'
    warn = '   [basis note]' if (sheet, lab) in BASIS_CHANGED else ''
    print(f'     {col:>3}  {lab:<38} {tag} {item}{warn}')

print(f'\n{"="*92}\nKEPT as cell-reference math ({len(kept)})\n{"="*92}')
for sheet, lab in kept:
    print(f'  {sheet} / {lab}')

print(f'\n{"="*92}\nNO CONSENSUS EQUIVALENT — left on their existing formula ({len(unmapped)})\n{"="*92}')
cur = None
for sheet, lab, f in unmapped:
    if sheet != cur:
        print(f'\n  {sheet}')
        cur = sheet
    print(f'     {lab}')

print(f'\n{"="*92}\nBUG FIX — lower-block period reference\n{"="*92}')
for sheet, lab, col in fixes:
    print(f'  {sheet} / {lab} at {col}: now reads its own period row 24 (was blank row 5)')
